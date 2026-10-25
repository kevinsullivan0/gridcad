"""HTTP regressions for the workbook-first input flow.

All workbooks are synthetic temporary files. Injected formula caches exercise the
reader contract, not an Excel calculation engine. These are API tests, not browser,
Windows, Excel, GrabCAD, or print-readiness validation.
"""
import hashlib
import json
from pathlib import Path
import tempfile
import threading
import unittest
from urllib.error import HTTPError
from urllib.request import Request, urlopen
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from grabcad_grid_ui import BUILD, ReviewServer


RECIPE = '20 / 0 / 0 / 0 / 30 / 50'
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
PROJECT = Path(__file__).parent / 'fixtures' / 'DO_NOT_PRINT_synthetic.PRINT'
DEFAULT_TOKEN = object()


def save_workbook(path, *, formula=False, scalar=False, multiple=False):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Recipes'
    ws['A1'] = 'Heading outside recipe selection'
    ws['A2'] = 20
    ws['B2'] = '=A2&" / 0 / 0 / 0 / 30 / 50"' if formula else 100 if scalar else RECIPE
    ws['C2'] = '0 / 10 / 0 / 0 / 30 / 60'
    ws['C3'] = '0;0;10;0;30;60'
    if multiple:
        settings = wb.create_sheet('Settings')
        settings['D900'] = 'A late sidebar value'
        wb.create_sheet('Empty')
        wb.active = 1
    wb.save(path)
    wb.close()


def inject_cache(path, value, cache_type='str', coordinate='B2'):
    """Change only a synthetic formula's persisted OOXML cached value."""
    with ZipFile(path) as archive:
        entries = [(entry, archive.read(entry.filename)) for entry in archive.infolist()]
    with ZipFile(path, 'w') as archive:
        for entry, payload in entries:
            if entry.filename == 'xl/worksheets/sheet1.xml':
                root = ET.fromstring(payload)
                cell = next(c for c in root.iter(NS + 'c') if c.attrib.get('r') == coordinate)
                cell.set('t', cache_type)
                saved = cell.find(NS + 'v')
                if saved is None:
                    saved = ET.SubElement(cell, NS + 'v')
                saved.text = value
                payload = ET.tostring(root, encoding='utf-8')
            archive.writestr(entry, payload)


class InputFlowAPITests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='gridcad-input-flow-test-')
        self.workbook = Path(self.tmp.name) / 'Synthetic workbook.xlsx'
        self.server = ReviewServer()
        self.worker = threading.Thread(
            target=self.server.serve_forever, kwargs={'poll_interval': 0.01}, daemon=True
        )
        self.worker.start()

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.worker.join(timeout=5)
        self.tmp.cleanup()

    def post(self, endpoint, data, *, binary=False, token=DEFAULT_TOKEN):
        headers = {'Content-Type': 'application/octet-stream' if binary else 'application/json'}
        if token is DEFAULT_TOKEN:
            token = self.server.token
        if token is not None:
            headers['X-GridCAD-Token'] = token
        request = Request(
            self.server.origin + endpoint,
            data=data if binary else json.dumps(data).encode('utf-8'),
            headers=headers,
            method='POST',
        )
        try:
            response = urlopen(request, timeout=10)
        except HTTPError as error:
            response = error
        with response:
            self.assertEqual(response.headers.get_content_type(), 'application/json')
            return response.status, json.load(response)

    def upload(self, **kwargs):
        save_workbook(self.workbook, **kwargs)
        return self.upload_existing()

    def upload_existing(self):
        status, value = self.post('/api/upload/workbook', self.workbook.read_bytes(), binary=True)
        self.assertEqual(status, 200, value)
        self.assertIs(value['source_modified'], False)
        return value

    def upload_project(self):
        status, value = self.post('/api/upload/project', PROJECT.read_bytes(), binary=True)
        self.assertEqual(status, 200, value)

    def check(self, **overrides):
        options = {'sheet': 'Recipes', 'range': 'B2:C3'}
        options.update(overrides)
        return self.post('/api/workbook-check', options)

    def assert_rejected_by_both(self, regions, expected):
        for endpoint in ('/api/workbook-check', '/api/analyze'):
            with self.subTest(endpoint=endpoint, regions=regions):
                status, value = self.post(endpoint, {
                    'sheet': 'Recipes', 'range': 'B2:C3', 'workbook_regions': regions,
                })
                self.assertEqual(status, 400, value)
                self.assertIn(expected, value['error'])
                self.assertIs(value['source_modified'], False)

    def test_session_is_ready_without_uploads_and_identifies_build(self):
        status, value = self.post('/api/session', {})
        self.assertEqual(status, 200)
        self.assertEqual(value['build'], BUILD)
        self.assertIs(value['session_ready'], True)
        self.assertEqual(value['max_recipe_cells'], 100000)
        self.assertIs(value['source_modified'], False)
        self.assertEqual(self.server.files, {})

    def test_session_rejects_bad_and_missing_tokens(self):
        for token in ('stale-token', '', None):
            with self.subTest(token=token):
                status, value = self.post('/api/session', {}, token=token)
                self.assertEqual(status, 403)
                self.assertIn('Session expired or missing', value['error'])
        self.assertEqual(self.post('/api/session', {})[0], 200)

    def test_session_rejects_unexpected_options(self):
        status, value = self.post('/api/session', {'sheet': 'Recipes'})
        self.assertEqual(status, 400)
        self.assertIn('no options', value['error'])

    def test_workbook_check_requires_workbook_but_not_project(self):
        status, value = self.check()
        self.assertEqual(status, 400)
        self.assertIn('Choose a workbook first', value['error'])
        self.upload()
        self.assertEqual(set(self.server.files), {'workbook'})
        status, value = self.check()
        self.assertEqual(status, 200, value)
        grid = value['grid']
        self.assertEqual((grid['rows'], grid['columns']), (2, 2))
        self.assertEqual([c['cell'] for c in grid['cells']], ['B2', 'C2', 'B3', 'C3'])
        self.assertEqual([c['role'] for c in grid['cells']], ['sample', 'sample', 'hole', 'sample'])
        self.assertEqual(grid['errors'], [])
        self.assertEqual(value['workbook_regions'], [])
        self.assertIs(value['source_modified'], False)
        self.assertIsNone(self.server.report)
        self.assertIsNone(self.server.analysis)

    def test_metadata_reports_active_tab_and_sparse_and_empty_origins(self):
        value = self.upload(multiple=True)
        info = value['workbook_info']
        self.assertEqual([sheet['name'] for sheet in info['sheets']], ['Recipes', 'Settings', 'Empty'])
        self.assertEqual(info['active_sheet'], 'Settings')
        self.assertEqual(info['preview_origins']['Settings'], 'D900')
        self.assertEqual(info['preview_origins']['Empty'], 'A1')
        status, preview = self.post('/api/workbook-preview', {
            'sheet': info['active_sheet'], 'start_cell': info['preview_origins'][info['active_sheet']],
        })
        self.assertEqual(status, 200, preview)
        self.assertEqual(preview['preview_cells'][0]['cell'], 'D900')
        self.assertEqual(preview['preview_cells'][0]['raw'], 'A late sidebar value')
        self.assertLessEqual(len(preview['preview_cells']), 1200)

    def test_sparse_origin_exposes_real_content_not_an_empty_bounding_corner(self):
        wb = Workbook()
        wb.active.title = 'Sparse'
        wb.active['B1000'] = RECIPE
        wb.active['AZ2'] = 'Early row, far-right content'
        wb.save(self.workbook)
        wb.close()
        info = self.upload_existing()['workbook_info']
        status, preview = self.post('/api/workbook-preview', {
            'sheet': 'Sparse', 'start_cell': info['preview_origins']['Sparse'],
        })
        self.assertEqual(status, 200, preview)
        self.assertTrue(
            any(cell['raw'] is not None for cell in preview['preview_cells']),
            'Recommended origin must reveal actual content, not the empty corner between B1000 and AZ2',
        )

    def test_multiple_tabs_require_explicit_check_sheet(self):
        self.upload(multiple=True)
        status, value = self.check(sheet='')
        self.assertEqual(status, 400)
        self.assertIn('Available tabs', value['error'])
        self.assertEqual(self.check()[0], 200)

    def test_invalid_preview_address_returns_json_then_valid_retry_works(self):
        self.upload()
        for address in ('A0', 'A1:B2', 'XFE1'):
            with self.subTest(address=address):
                status, value = self.post('/api/workbook-preview', {'sheet': 'Recipes', 'start_cell': address})
                self.assertEqual(status, 400, value)
                self.assertIn('Preview starts', value['error'])
        status, value = self.post('/api/workbook-preview', {'sheet': 'Recipes', 'start_cell': 'B2'})
        self.assertEqual(status, 200, value)
        self.assertEqual(value['preview_cells'][0]['raw'], RECIPE)

    def test_malformed_shared_string_cache_falls_back_to_raw_cells_without_admitting_formulas(self):
        save_workbook(self.workbook, formula=True)
        inject_cache(self.workbook, '999', cache_type='s')
        self.upload_existing()  # Formula metadata is readable despite the broken cached value.
        status, value = self.post('/api/workbook-preview', {'sheet': 'Recipes', 'start_cell': 'B2'})
        self.assertEqual(status, 200, value)
        self.assertIn('Saved formula results could not be read', value['cache_read_error'])
        self.assertTrue(value['preview_cells'][0]['raw'].startswith('='))
        self.assertIsNone(value['preview_cells'][0]['cached_value'])
        self.assertEqual(self.check(formula_cache_confirmed=True)[0], 400)
        self.assertIs(value['source_modified'], False)
        self.assertEqual(self.post('/api/session', {})[0], 200)
        # A literal-only range remains checkable without reading that unrelated cache.
        status, value = self.check(range='C2')
        self.assertEqual(status, 200, value)
        self.assertEqual(value['grid']['cells'][0]['role'], 'sample')

    @unittest.skipUnless(Path('/proc/self/fd').exists(), 'Linux file-descriptor lifecycle check')
    def test_malformed_cache_closes_input_stream_without_waiting_for_gc(self):
        import gc
        from gridcad.core import read_workbook
        save_workbook(self.workbook, formula=True)
        inject_cache(self.workbook, '999', cache_type='s')
        gc.collect()
        enabled=gc.isenabled();gc.disable()
        before=len(list(Path('/proc/self/fd').iterdir()))
        try:
            for _ in range(5):
                preview=read_workbook(self.workbook,metadata_only=True,sheet='Recipes',cell_range='C2')
                self.assertIsNotNone(preview['cache_read_error'])
            after=len(list(Path('/proc/self/fd').iterdir()))
            self.assertLessEqual(after,before)
        finally:
            if enabled:gc.enable()
            gc.collect()

    def test_nonfinite_saved_cache_never_emits_invalid_json(self):
        import json
        save_workbook(self.workbook, formula=True)
        inject_cache(self.workbook, '1e999', cache_type='n')
        self.upload_existing()
        status, preview=self.post('/api/workbook-preview',{'sheet':'Recipes','start_cell':'A1'})
        self.assertEqual(status,200)
        from gridcad.core import json_text
        json.loads(json_text(preview),parse_constant=lambda value: (_ for _ in ()).throw(AssertionError(value)))
        status,checked=self.check(formula_cache_confirmed=True)
        self.assertEqual(status,200)
        self.assertEqual(checked['grid']['cells'][0]['formula_status'],'nonfinite_saved_result')
        self.assertEqual(checked['grid']['cells'][0]['role'],'invalid')
        json.loads(json_text(checked),parse_constant=lambda value: (_ for _ in ()).throw(AssertionError(value)))

    def test_literal_scalar_is_invalid_and_never_expanded_to_six_channels(self):
        self.upload(scalar=True)
        status, value = self.check()
        self.assertEqual(status, 200, value)
        scalar = value['grid']['cells'][0]
        self.assertEqual(scalar['role'], 'invalid')
        self.assertIsNone(scalar['recipe'])
        self.assertIn('not a scalar', scalar['error'])

    def test_saved_formula_requires_confirmation_then_preserves_exact_recipe(self):
        save_workbook(self.workbook, formula=True)
        inject_cache(self.workbook, RECIPE)
        original_hash = hashlib.sha256(self.workbook.read_bytes()).hexdigest()
        self.upload_existing()
        status, value = self.check()
        self.assertEqual(status, 200, value)
        formula = value['grid']['cells'][0]
        self.assertEqual(formula['role'], 'invalid')
        self.assertEqual(formula['formula_status'], 'confirmation_required')
        self.assertIsNone(formula['value_used'])
        status, value = self.check(formula_cache_confirmed=True, transpose=True)
        self.assertEqual(status, 200, value)
        formula = value['grid']['cells'][0]
        self.assertEqual(formula['role'], 'sample')
        self.assertEqual(formula['formula_status'], 'saved_result_confirmed')
        self.assertEqual(formula['raw'], '=A2&" / 0 / 0 / 0 / 30 / 50"')
        self.assertEqual(formula['cached_value'], RECIPE)
        self.assertEqual(formula['recipe']['total'], '100')
        self.assertEqual(formula['recipe']['channels'], dict(zip('CMYKWT', ['20', '0', '0', '0', '30', '50'])))
        self.assertEqual(formula['cache_freshness'], 'not_verifiable')
        self.assertEqual(hashlib.sha256(self.workbook.read_bytes()).hexdigest(), original_hash)

    def test_non_boolean_formula_confirmation_does_not_authorize_cache(self):
        save_workbook(self.workbook, formula=True)
        inject_cache(self.workbook, RECIPE)
        self.upload_existing()
        for confirmation in ('true', 1, None):
            with self.subTest(confirmation=confirmation):
                status, value = self.check(formula_cache_confirmed=confirmation)
                self.assertEqual(status, 200, value)
                self.assertEqual(value['grid']['cells'][0]['formula_status'], 'confirmation_required')

    def test_missing_formula_cache_is_invalid_not_a_hole_even_if_confirmed(self):
        self.upload(formula=True)
        for confirmation in (False, True):
            with self.subTest(confirmation=confirmation):
                status, value = self.check(formula_cache_confirmed=confirmation)
                self.assertEqual(status, 200, value)
                formula = value['grid']['cells'][0]
                self.assertEqual(formula['role'], 'invalid')
                self.assertEqual(formula['formula_status'], 'missing_saved_result')
                self.assertEqual(value['grid']['cells'][2]['role'], 'hole')

    def test_confirmed_scalar_formula_cache_stays_invalid(self):
        save_workbook(self.workbook, formula=True)
        inject_cache(self.workbook, '100', cache_type='n')
        self.upload_existing()
        status, value = self.check(formula_cache_confirmed=True)
        self.assertEqual(status, 200, value)
        formula = value['grid']['cells'][0]
        self.assertEqual(formula['role'], 'invalid')
        self.assertEqual(formula['cached_value'], 100)
        self.assertIn('not a scalar', formula['error'])

    def test_saved_excel_error_is_not_reported_as_missing_cache(self):
        save_workbook(self.workbook, formula=True)
        inject_cache(self.workbook, '#VALUE!', cache_type='e')
        self.upload_existing()
        status, value = self.check(formula_cache_confirmed=True)
        self.assertEqual(status, 200, value)
        formula = value['grid']['cells'][0]
        self.assertEqual(formula['role'], 'invalid')
        self.assertEqual(formula['formula_status'], 'saved_excel_error')
        self.assertIn('#VALUE!', formula['error'])

    def test_array_and_data_table_regions_block_even_with_anchor_outside_selection(self):
        for structured in (
            ArrayFormula(ref='A1:B1', text='=IF({1,2},"20/0/0/0/30/50")'),
            DataTableFormula(ref='A1:B1', r1='D1'),
        ):
            with self.subTest(structured=type(structured).__name__):
                wb = Workbook()
                wb.active.title = 'Recipes'
                wb.active['A1'] = structured
                wb.active['B1'] = RECIPE
                wb.active['B2'] = RECIPE
                wb.save(self.workbook)
                wb.close()
                self.upload_existing()
                status, preview = self.post('/api/workbook-preview', {'sheet': 'Recipes', 'start_cell': 'B1'})
                self.assertEqual(status, 200, preview)
                self.assertEqual(preview['preview_cells'][0]['unsupported_formula_region'], 'A1 (A1:B1)')
                status, value = self.check(range='B1', formula_cache_confirmed=True)
                self.assertEqual(status, 400, value)
                self.assertIn('array/spill or data-table', value['error'])
                self.assertEqual(self.check(range='B2')[0], 200)

    def test_nonoverlapping_reference_regions_preserve_selection_in_check_and_analyze(self):
        self.upload()
        self.upload_project()
        regions = [{'role': 'header', 'range': 'B1:C1'}, {'role': 'sidebar', 'range': 'A2:A3'}]
        for endpoint in ('/api/workbook-check', '/api/analyze'):
            with self.subTest(endpoint=endpoint):
                status, value = self.post(endpoint, {
                    'sheet': 'Recipes', 'range': 'B2:C3', 'workbook_regions': regions,
                })
                self.assertEqual(status, 200, value)
                self.assertEqual(value['workbook_regions'], regions)
                self.assertEqual(value['grid']['range'], 'B2:C3')
                self.assertEqual(len(value['grid']['cells']), 4)
                self.assertEqual(value['grid']['cells'][2]['role'], 'hole')

    def test_recipe_annotation_overlap_is_rejected_by_check_and_analyze(self):
        self.upload()
        self.upload_project()
        for role, rectangle in (('header', 'B1:B2'), ('sidebar', 'C3:D4'), ('header', 'B2:C3')):
            self.assert_rejected_by_both([{'role': role, 'range': rectangle}], 'Recipe selection overlaps')

    def test_annotation_annotation_overlap_is_rejected_by_check_and_analyze(self):
        self.upload()
        self.upload_project()
        self.assert_rejected_by_both([
            {'role': 'header', 'range': 'E1:F2'}, {'role': 'sidebar', 'range': 'F2:G3'},
        ], 'reference regions overlap')

    def test_reference_limit_accepts_32_and_rejects_33_in_both_endpoints(self):
        self.upload()
        self.upload_project()
        regions = [{'role': 'sidebar', 'range': f'Z{row}'} for row in range(1, 34)]
        for endpoint in ('/api/workbook-check', '/api/analyze'):
            with self.subTest(endpoint=endpoint):
                status, value = self.post(endpoint, {
                    'sheet': 'Recipes', 'range': 'B2:C3', 'workbook_regions': regions[:32],
                })
                self.assertEqual(status, 200, value)
                self.assertEqual(len(value['workbook_regions']), 32)
        self.assert_rejected_by_both(regions, 'at most 32')

    def test_reference_regions_reject_invalid_shapes_roles_and_ranges(self):
        self.upload()
        self.upload_project()
        for regions, diagnostic in (
            ({'role': 'header', 'range': 'A1'}, 'at most 32'),
            ([{'role': 'sample', 'range': 'A1'}], 'sidebar or header'),
            ([{'role': 'header', 'range': 'A1', 'ignore': True}], 'sidebar or header'),
            ([{'role': 'header', 'range': 'A:A'}], 'finite Excel cell rectangle'),
            ([{'role': 'header', 'range': 'XFE1'}], 'finite Excel cell rectangle'),
            ([{'role': 'header', 'range': 12}], 'Excel cell rectangle'),
        ):
            self.assert_rejected_by_both(regions, diagnostic)

    def test_workbook_check_revokes_prior_mapping_state(self):
        self.upload()
        self.upload_project()
        status, value = self.post('/api/analyze', {'sheet': 'Recipes', 'range': 'B2:C3'})
        self.assertEqual(status, 200, value)
        self.assertIsNotNone(self.server.analysis)
        self.assertIsNotNone(self.server.report)
        self.assertEqual(self.check()[0], 200)
        self.assertIsNone(self.server.analysis)
        self.assertIsNone(self.server.report)
        self.assertIsNone(self.server.review)


if __name__ == '__main__':
    unittest.main()
