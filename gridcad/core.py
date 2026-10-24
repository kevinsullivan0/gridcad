from __future__ import annotations
import csv
import hashlib
import html
import json
import re
import stat
import unicodedata
import zipfile
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path, PurePosixPath

CHANNELS = 'CMYKWT'
FIELDS = ('cyan', 'magenta', 'yellow', 'black', 'white', 'clear')
MAX_ARCHIVE = 2 * 1024**3
MAX_ENTRY = 512 * 1024**2
MAX_TOTAL = 2 * 1024**3
MAX_JSON = 32 * 1024**2
MAX_ENTRIES = 20000
MAX_CELLS = 100000

class ValidationError(ValueError):
    pass

def sha256_bytes(data):
    return hashlib.sha256(data).hexdigest()

def file_hash(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for b in iter(lambda: f.read(1024 * 1024), b''):
            h.update(b)
    return h.hexdigest()

def _pairs(pairs):
    out = {}
    for k, v in pairs:
        if k in out:
            raise ValidationError(f'Duplicate JSON key: {k}')
        out[k] = v
    return out

def strict_json(data):
    try:
        value = json.loads(data, object_pairs_hook=_pairs, parse_float=Decimal,
                          parse_constant=lambda s: (_ for _ in ()).throw(ValidationError(f'Nonfinite JSON: {s}')))
        stack=[(value,0)]
        while stack:
            item,depth=stack.pop()
            if depth>64: raise ValidationError('JSON nesting exceeds 64 levels.')
            if isinstance(item,dict): stack.extend((x,depth+1) for x in item.values())
            elif isinstance(item,list): stack.extend((x,depth+1) for x in item)
        return value
    except (ValueError, UnicodeError, RecursionError) as exc:
        raise ValidationError(f'Invalid JSON: {exc}') from exc

def json_text(value):
    # Reports serialize Decimal as an exact decimal string, never a rounded float.
    return json.dumps(value, ensure_ascii=False, indent=2, default=lambda x: str(x) if isinstance(x, Decimal) else (_ for _ in ()).throw(TypeError(type(x).__name__)))

def validate_recipe(values, *, order='CMYKWT', tolerance=Decimal(0)):
    if len(order) != 6 or set(order) != set(CHANNELS):
        raise ValidationError('Channel order must explicitly contain C M Y K W T exactly once.')
    if not tolerance.is_finite() or tolerance < 0 or tolerance > Decimal('0.1'):
        raise ValidationError('Sum tolerance must be finite and between 0 and 0.1 percentage points.')
    if len(values) != 6:
        raise ValidationError('Exactly six positional channel values are required; missing positions are ambiguous.')
    numbers = []
    for value in values:
        if isinstance(value, bool):
            raise ValidationError('Boolean is not a channel percentage.')
        try:
            d = Decimal(str(value).strip())
        except (InvalidOperation, ValueError):
            raise ValidationError(f'Invalid channel percentage: {value!r}')
        if d.is_finite() and (len(d.as_tuple().digits)>64 or abs(d.as_tuple().exponent)>64):
            raise ValidationError('Channel decimal exceeds supported 64-digit/exponent precision limit.')
        if not d.is_finite() or not Decimal(0) <= d <= Decimal(100):
            raise ValidationError(f'Channel percentage outside finite 0..100: {value!r}')
        numbers.append(d)
    with localcontext() as ctx:
        ctx.prec = max(64, sum(len(x.as_tuple().digits) + abs(x.as_tuple().exponent) for x in numbers) + 10)
        total = sum(numbers, Decimal(0))
        difference = abs(total - Decimal(100))
    if difference > tolerance:
        raise ValidationError(f'Channel total is {total}, not 100 (tolerance {tolerance}). No normalization performed.')
    recipe = dict(zip(order, numbers))
    return {'channels': {k: str(recipe[k]) for k in CHANNELS}, 'total': str(total), 'sum_tolerance': str(tolerance)}

def parse_bundle(value, *, order='CMYKWT', empty_channels_zero=False, tolerance=Decimal(0), delimiter='auto'):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None  # A grid hole, never an all-zero recipe.
    if not isinstance(value, str):
        raise ValidationError('A recipe cell must contain a six-number bundle, not a scalar.')
    s = value.strip()
    if unicodedata.normalize('NFKC', s) != s:
        raise ValidationError('Compatibility Unicode characters need explicit conversion; no silent normalization.')
    if len(s) > 512:
        raise ValidationError('Recipe cell is too long.')
    if s.startswith('='):
        raise ValidationError('Formula cells are not evaluated; supply literal recipe values.')
    if (s.startswith('[') and s.endswith(']')) or (s.startswith('(') and s.endswith(')')):
        s = s[1:-1].strip()
    if delimiter == 'auto':
        present = [d for d in ('/', ';', ',') if d in s]
        if len(present) > 1:
            raise ValidationError('Mixed bundle separators are ambiguous; choose a consistent schema.')
        if ',' in present: raise ValidationError('Comma bundles require explicit --delimiter comma; decimal-comma ambiguity cannot be inferred.')
        delim = present[0] if present else None
    else:
        delim = {'slash': '/', 'comma': ',', 'semicolon': ';', 'space': None}.get(delimiter)
        if delimiter not in ('slash', 'comma', 'semicolon', 'space'):
            raise ValidationError('Unknown delimiter.')
    parts = s.split(delim) if delim else s.split()
    if len(parts) != 6:
        raise ValidationError('Expected six positions. Five values cannot identify the missing channel.')
    tokens = []
    for p in parts:
        p = p.strip()
        if not p:
            if not empty_channels_zero:
                raise ValidationError('Empty channel position; use the explicit empty-channel-zero option if this schema means zero.')
            p = '0'
        if p.endswith('%'):
            p = p[:-1].strip()
        if not re.fullmatch(r'[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)', p):
            raise ValidationError(f'Unsupported numeric token {p!r}; decimal point and plain decimal numbers required.')
        tokens.append(p)
    return validate_recipe(tokens, order=order, tolerance=tolerance)

def archive_inventory(path, *, require_project=True):
    path = Path(path)
    if path.stat().st_size > MAX_ARCHIVE:
        raise ValidationError('Archive exceeds compressed size limit.')
    entries, folded, names = [], set(), set()
    total = 0
    project_data = None
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
            if len(infos) > MAX_ENTRIES:
                raise ValidationError('Too many archive entries.')
            for info in infos:
                name = info.filename
                if info.orig_filename != name:
                    raise ValidationError('Archive filename was truncated or normalized by ZIP reader.')
                posix = PurePosixPath(name)
                if ('\\' in name or '//' in name or '\x00' in name or ':' in name or name.startswith('/') or
                    any(x in ('..', '.') for x in name.split('/')) or not name or
                    any(ord(c) < 32 or c in '<>"|?*' for c in name)):
                    raise ValidationError(f'Unsafe archive path: {name!r}')
                for part in name.rstrip('/').split('/'):
                    stem=part.split('.')[0].upper()
                    if part.endswith((' ','.')) or stem in {'CON','PRN','AUX','NUL',*[f'COM{i}' for i in range(1,10)],*[f'LPT{i}' for i in range(1,10)]}:
                        raise ValidationError('Archive path is not safe on Windows: '+name)
                normalized = unicodedata.normalize('NFC', name).casefold().rstrip('/')
                if name in names or normalized in folded:
                    raise ValidationError(f'Duplicate/case-colliding archive entry: {name}')
                names.add(name); folded.add(normalized)
                mode = info.external_attr >> 16
                if stat.S_ISLNK(mode):
                    raise ValidationError(f'Symlink archive entry: {name}')
                if info.flag_bits & 1:
                    raise ValidationError('Encrypted archives are not supported.')
                if info.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
                    raise ValidationError(f'Unsupported ZIP compression for {name}')
                total += info.file_size
                if info.file_size > MAX_ENTRY or total > MAX_TOTAL:
                    raise ValidationError('Archive expansion exceeds safety limit.')
                if info.file_size > 1024**2 and info.file_size / max(info.compress_size, 1) > 200:
                    raise ValidationError(f'Excessive compression ratio: {name}')
                if name == 'project.json' and info.file_size > MAX_JSON:
                    raise ValidationError('project.json exceeds size limit.')
            # No extraction. Hash streamed payloads and verify CRC through ZipFile.
            for info in infos:
                h = hashlib.sha256(); count = 0; chunks = []
                with z.open(info) as f:
                    while True:
                        b = f.read(1024 * 1024)
                        if not b:
                            break
                        count += len(b)
                        if count > info.file_size or count > MAX_ENTRY:
                            raise ValidationError('Expanded entry exceeds declared limit.')
                        h.update(b)
                        if info.filename == 'project.json':
                            chunks.append(b)
                if count != info.file_size:
                    raise ValidationError('Entry length differs from declaration.')
                entries.append({'name': info.filename, 'size': count, 'sha256': h.hexdigest(), 'directory': info.is_dir()})
                if info.filename == 'project.json':
                    project_data = b''.join(chunks)
    except (zipfile.BadZipFile, RuntimeError, OSError, NotImplementedError) as exc:
        raise ValidationError(f'Unreadable archive: {exc}') from exc
    if require_project and project_data is None:
        raise ValidationError('Root project.json is missing.')
    return {'path': str(path), 'sha256': file_hash(path), 'entries': entries,
            'expanded_bytes': total, 'project_document': strict_json(project_data) if project_data is not None else None}

def _list(value, field):
    if not isinstance(value, list):
        raise ValidationError(f'{field} must be an array; unknown layout refused.')
    return value

def project_index(document):
    if not isinstance(document, dict) or not isinstance(document.get('project'), dict) or 'version' not in document:
        raise ValidationError('Expected version/project wrapper; unknown layout refused.')
    p = document['project']; occurrences = {}; body_refs = {}; tray_ids = set(); body_ids = set()
    def visit(nodes, depth=0):
        if depth > 64:
            raise ValidationError('Occurrence hierarchy is too deep.')
        for node in _list(nodes, 'occurrences'):
            if not isinstance(node, dict) or not isinstance(node.get('id'), str) or not node['id']:
                raise ValidationError('Occurrence requires a nonempty stable string ID.')
            oid = node['id']
            if oid in occurrences:
                raise ValidationError(f'Duplicate occurrence ID: {oid}')
            occurrences[oid] = node
            visit(node.get('children', []), depth + 1)
    visit(p.get('rootOccurrences'))
    for tray in _list(p.get('trays'), 'trays'):
        if not isinstance(tray, dict) or not isinstance(tray.get('id'), str) or not tray['id'] or tray['id'] in tray_ids:
            raise ValidationError('Missing or duplicate tray ID.')
        tray_ids.add(tray['id'])
        for body in _list(tray.get('trayBodies'), 'trayBodies'):
            if not isinstance(body, dict) or not isinstance(body.get('occurrence'), str):
                raise ValidationError('Tray body requires occurrence reference.')
            if body['occurrence'] not in occurrences:
                raise ValidationError(f'Dangling tray body occurrence: {body["occurrence"]}')
            if not isinstance(body.get('id'), str) or not body['id'] or body['id'] in body_ids:
                raise ValidationError('Missing or duplicate tray body ID.')
            body_ids.add(body['id'])
            body_refs.setdefault(body['occurrence'], []).append((tray, body))
    return p, occurrences, body_refs

def material_mode(material):
    """Conservative inspection of observed 2026 vendor discriminator precedence."""
    if not isinstance(material, dict): return 'unknown'
    m=material.get('metadata') or {}
    if not isinstance(m,dict): return 'unknown'
    if m.get('materialType'): return m['materialType']
    if str(material.get('name','')).startswith('rgba') or m.get('pickerSelectedColor'): return 'color_picker'
    if m.get('pantone_materials'): return 'pantone'
    if m.get('cmykwt_dm_used_materials_id'): return 'cmykwt_dm_mixer'
    if m.get('dmCategoryName'): return 'digital_material'
    return 'base_material'

def material_summary(material):
    if not isinstance(material, dict):
        return {'recipe': None, 'resin_ids': [], 'issue': 'No material object.'}
    metadata = material.get('metadata')
    if not isinstance(metadata, dict):
        return {'recipe': None, 'resin_ids': [], 'issue': 'No CMYKWT metadata.'}
    if material_mode(material) == 'base_material':
        return {'recipe':None,'resin_ids':[], 'issue':None,'name':material.get('name'), 'vendor_material_id':material.get('vendorMaterialId'),'resin_id':metadata.get('resinId'),'mode':'base_material'}
    values = [metadata.get('cmykwt_dm_' + f) for f in FIELDS]
    try:
        recipe = validate_recipe(values)
        issue = None
    except ValidationError as exc:
        recipe = None; issue = str(exc)
    ids = metadata.get('cmykwt_dm_used_materials_id')
    if not isinstance(ids, list) or not ids or any(not ((isinstance(x,str) and x) or (type(x) is int and x>0)) for x in ids):
        issue = (issue + ' ' if issue else '') + 'Missing/unknown actual resin identity list.'
        ids = []
    return {'recipe': recipe, 'resin_ids': [str(x) for x in ids], 'issue': issue,
            'name': material.get('name'), 'vendor_material_id': material.get('vendorMaterialId'), 'mode': material_mode(material)}

def inspect_project(path):
    inv = archive_inventory(path); p, occ, bodies = project_index(inv['project_document'])
    return {'status': 'READ_ONLY_INSPECTION', 'application_validated': False,
            'version': inv['project_document']['version'], 'project_id': p.get('id'),
            'template': p.get('manufacturingTemplateName'), 'archive': {k: v for k, v in inv.items() if k != 'project_document'},
            'occurrences': [{'id': oid, 'name': n.get('name'), 'children': [c['id'] for c in n.get('children', [])],
                             'transform': n.get('transform'), 'desired_material': material_summary(n.get('material')),
                             'print_bodies': [{'tray_id': t['id'], 'body_id': b['id'], 'transform': b.get('transform'),
                                               'compose_transform': b.get('composeTransform'),
                                               'actual_material': material_summary(b.get('material'))} for t,b in bodies.get(oid, [])]}
                            for oid,n in occ.items()],
            'warning': 'Read-only observed JSON structure. This inspection changes nothing and does not establish GUI reopen, composition or print readiness. A separate candidate-copy gate is required before any save.'}

def read_workbook(path, *, sheet=None, cell_range=None, formula_cache_confirmed=False, metadata_only=False, **parse_options):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    path = Path(path)
    if path.suffix.lower() != '.xlsx':
        raise ValidationError('Only .xlsx is supported. Macros and legacy binary workbooks are not executed.')
    inv=archive_inventory(path, require_project=False)  # Bound decompression before XML parsing.
    if inv['expanded_bytes']>64*1024**2: raise ValidationError('Workbook exceeds 64 MiB expanded preview limit.')
    # Resolve XML by OPC content types and relationships, never suffix alone.
    from defusedxml import ElementTree as ET
    from defusedxml.common import DefusedXmlException
    from io import BytesIO
    import posixpath
    allocated_cells=0
    with zipfile.ZipFile(path) as z:
        entries={entry['name']:entry for entry in inv['entries']}
        def xml_bytes(name):
            if name not in entries: raise ValidationError('Workbook XML relationship target is missing: '+name)
            if entries[name]['size']>16*1024**2:
                raise ValidationError('Workbook XML part exceeds 16 MiB preview limit.')
            raw=z.read(name)
            if b'\x00' in raw: raise ValidationError('Workbook XML must use UTF-8-compatible encoding.')
            if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():
                raise ValidationError('Workbook XML DTD/entities are not supported.')
            return raw
        def parse_xml(name):
            try: return ET.fromstring(xml_bytes(name), forbid_dtd=True, forbid_entities=True, forbid_external=True)
            except (ET.ParseError,DefusedXmlException) as exc: raise ValidationError('Malformed or unsafe workbook XML: '+name) from exc
        manifest=parse_xml('[Content_Types].xml')
        defaults={}; overrides={}
        for elem in manifest:
            kind=elem.tag.rsplit('}',1)[-1]
            if kind=='Default':
                key=elem.attrib.get('Extension','').lower()
                if key in defaults: raise ValidationError('Duplicate workbook content-type Default declaration.')
                defaults[key]=elem.attrib.get('ContentType','')
            elif kind=='Override':
                key=elem.attrib.get('PartName','').lstrip('/')
                if key in overrides: raise ValidationError('Duplicate workbook content-type Override declaration.')
                overrides[key]=elem.attrib.get('ContentType','')
        xml_parts={'[Content_Types].xml'}
        for name in entries:
            mime=overrides.get(name,defaults.get(name.rsplit('.',1)[-1].lower(),''))
            if name.lower().endswith(('.xml','.rels')) or 'xml' in mime.lower(): xml_parts.add(name)
        # A worksheet relationship requires XML regardless of its extension or MIME.
        # Other XML-bearing Office relationships are covered too; image/media payloads
        # remain binary and are not fed to an XML parser.
        xml_relation_kinds={'worksheet','chartsheet','dialogsheet','macrosheet','officeDocument',
            'styles','sharedStrings','theme','comments','table','drawing','chart','externalLink',
            'pivotTable','pivotCacheDefinition','pivotCacheRecords','calcChain','vmlDrawing'}
        for name in list(xml_parts):
            if not name.lower().endswith('.rels'): continue
            root=parse_xml(name)
            base='' if name=='_rels/.rels' else posixpath.dirname(posixpath.dirname(name))
            for rel in root:
                if rel.attrib.get('TargetMode')=='External': continue
                if rel.attrib.get('Type','').rsplit('/',1)[-1] not in xml_relation_kinds: continue
                target=rel.attrib.get('Target','')
                resolved=posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join(base,target))
                if resolved.startswith('../') or resolved=='.' or '\\' in resolved:
                    raise ValidationError('Invalid workbook XML relationship target.')
                xml_parts.add(resolved)
        for name in sorted(xml_parts):
            try:
                for event,elem in ET.iterparse(BytesIO(xml_bytes(name)),events=('end',),
                        forbid_dtd=True,forbid_entities=True,forbid_external=True):
                    tag=elem.tag.rsplit('}',1)[-1]
                    if tag=='c': allocated_cells+=1
                    if tag=='mergeCell':
                        try: mc,mr,xc,xr=range_boundaries(elem.attrib['ref'])
                        except (ValueError,KeyError): raise ValidationError('Malformed merged range.')
                        if None in (mc,mr,xc,xr) or xc<mc or xr<mr: raise ValidationError('Invalid merged range.')
                        allocated_cells+=(xc-mc+1)*(xr-mr+1)
                    if allocated_cells>MAX_CELLS: raise ValidationError('Workbook physical/merged cell allocation exceeds limit.')
                    elem.clear()
            except (ET.ParseError,DefusedXmlException) as exc: raise ValidationError('Malformed or unsafe workbook XML: '+name) from exc
    try:
        wb = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValidationError('Workbook could not be loaded: '+str(exc)[:500]) from exc
    cached_wb = None
    try:
        sheets = [{'name': ws.title, 'used_range': ws.calculate_dimension()} for ws in wb.worksheets]
        def structured_regions(ws):
            regions=[]
            # Iterate physically allocated cells, not the potentially enormous
            # used rectangle. openpyxl is pinned; its formula objects expose ref.
            for cell in list(ws._cells.values()):
                if cell.data_type != 'f' or isinstance(cell.value, str): continue
                anchor=cell.coordinate
                ref=getattr(cell.value, 'ref', None) or anchor
                try: bounds=range_boundaries(ref)
                except (ValueError, TypeError) as exc: raise ValidationError('Unsupported structured formula range. Paste evaluated recipes as values in a workbook copy.') from exc
                if None in bounds: raise ValidationError('Unsupported unbounded structured formula range.')
                regions.append((anchor,ref,bounds))
            return regions
        if metadata_only:
            result = {'sheets': sheets, 'source_modified': False}
            if sheet is None: return result
            if sheet not in wb.sheetnames:
                raise ValidationError('Choose one of the available Excel tabs.')
            ws = wb[sheet]
            try: startc, startr, endc, endr = range_boundaries(cell_range or 'A1')
            except ValueError as exc: raise ValidationError('Preview starts at a cell address, for example A1.') from exc
            if None in (startc, startr, endc, endr) or startc != endc or startr != endr or not (1 <= startc <= 16384 and 1 <= startr <= 1048576):
                raise ValidationError('Preview starts at one Excel cell, for example A1.')
            # Bounded data grid, not a renderer or formula engine. Sidebar cells
            # are shown, never parsed as recipes until explicitly selected.
            endc = min(16384, max(startc, min(ws.max_column, startc+29)))
            endr = min(1048576, max(startr, min(ws.max_row, startr+39)))
            cached_wb = load_workbook(path, read_only=False, data_only=True, keep_links=False)
            arrays=structured_regions(ws)
            preview=[]
            for row in ws.iter_rows(min_row=startr,max_row=endr,min_col=startc,max_col=endc):
                for cell in row:
                    raw=cell.value
                    raw=raw if raw is None or type(raw) in (str,int,float,bool) else str(getattr(raw,'text',raw))
                    cached=cached_wb[sheet][cell.coordinate].value if cell.data_type=='f' else None
                    array_region=next((f'{anchor} ({ref})' for anchor,ref,(ac,ar,bc,br) in arrays if ac<=cell.column<=bc and ar<=cell.row<=br),None)
                    preview.append({'unsupported_formula_region':array_region,'cell':cell.coordinate,'row':cell.row,'column':cell.column,
                                    'raw':raw,'is_formula':cell.data_type=='f',
                                    'cached_value':cached if cached is None or type(cached) in (str,int,float,bool) else str(cached)})
            result.update(sheet=sheet, preview_cells=preview,
                          preview_range=f'{get_column_letter(startc)}{startr}:{get_column_letter(endc)}{endr}',
                          start_row=startr,start_column=startc,rows=endr-startr+1,columns=endc-startc+1,
                          column_labels=[get_column_letter(c) for c in range(startc,endc+1)],
                          preview_limited=ws.max_row>endr or ws.max_column>endc or startr>1 or startc>1)
            return result
        if sheet is None:
            if len(wb.sheetnames) != 1:
                raise ValidationError('Choose the Excel tab containing your recipe grid. Available tabs: ' + ', '.join(wb.sheetnames))
            sheet = wb.sheetnames[0]
        if sheet not in wb.sheetnames:
            raise ValidationError(f'Excel tab not found: {sheet}. Available tabs: ' + ', '.join(wb.sheetnames))
        ws = wb[sheet]
        bounds = range_boundaries(cell_range or ws.calculate_dimension())
        minc,minr,maxc,maxr = bounds
        if any(v is None for v in bounds) or minr < 1 or minc < 1 or maxr < minr or maxc < minc or maxr > 1048576 or maxc > 16384:
            raise ValidationError('Range must be a finite cell rectangle, e.g. B2:F4.')
        if (maxr-minr+1)*(maxc-minc+1) > MAX_CELLS:
            raise ValidationError('Selected range exceeds cell limit.')
        for merged in ws.merged_cells.ranges:
            if not (merged.max_row < minr or merged.min_row > maxr or merged.max_col < minc or merged.min_col > maxc):
                raise ValidationError(f'Merged cells intersect selected grid: {merged}; choose a grid without merged cells.')
        for anchor, ref, (ac, ar, bc, br) in structured_regions(ws):
            if not (br < minr or ar > maxr or bc < minc or ac > maxc):
                raise ValidationError(f'Selected recipe grid intersects unsupported array/spill or data-table formula {anchor} ({ref}). Structured formula ranges are not supported, even when their anchor is outside the selection. Paste evaluated recipes as values in a workbook copy, then select that copy.')
        cells=[]; errors=[]; formula_count=0
        # Read saved results separately. Neither openpyxl read executes formulas,
        # macros, links or a spreadsheet calculation engine. Trust is explicit.
        if any(cell.data_type == 'f' for row in ws.iter_rows(min_row=minr,max_row=maxr,min_col=minc,max_col=maxc) for cell in row):
            try:
                cached_wb = load_workbook(path, read_only=False, data_only=True, keep_links=False)
            except Exception as exc:
                raise ValidationError('Saved formula results could not be read: ' + str(exc)[:500]) from exc
        def display_value(value):
            return value if value is None or type(value) in (str,int,float,bool) else str(getattr(value, 'text', value))
        for row in ws.iter_rows(min_row=minr,max_row=maxr,min_col=minc,max_col=maxc):
            for cell in row:
                record={'cell':cell.coordinate,'row':cell.row,'column':cell.column,
                        'grid_row':cell.row-minr,'grid_column':cell.column-minc,
                        'raw':display_value(cell.value), 'raw_type':type(cell.value).__name__,
                        'value_source':'literal', 'value_used':display_value(cell.value)}
                try:
                    value = cell.value
                    if cell.data_type == 'f':
                        if not isinstance(cell.value, str):
                            raise ValidationError('This structured Excel formula type is not supported. Paste evaluated recipes as values in a workbook copy.')
                        formula_count += 1
                        cached_cell = cached_wb[sheet][cell.coordinate]
                        record.update(value_source='saved_formula_result', formula=display_value(cell.value),
                                      cached_value=display_value(cached_cell.value), value_used=None,
                                      cache_freshness='not_verifiable', cache_confirmed=formula_cache_confirmed is True)
                        if cached_cell.value is None:
                            record['formula_status']='missing_saved_result'
                            raise ValidationError('Formula has no readable saved result. Open the workbook in Excel, recalculate, save, then select the saved file again; or paste the evaluated recipe as values in a workbook copy. GridCAD does not calculate formulas. A missing result is not treated as a blank grid cell.')
                        if cached_cell.data_type == 'e':
                            record['formula_status']='saved_excel_error'
                            raise ValidationError('Saved formula result is an Excel error (' + str(cached_cell.value) + '). Fix it in Excel, recalculate and save, then select the file again.')
                        if formula_cache_confirmed is not True:
                            record['formula_status']='confirmation_required'
                            raise ValidationError('A saved formula result is available but its freshness cannot be verified. Recalculate and save in Excel, reselect the file, then confirm use of saved formula results; or paste evaluated recipes as values in a copy.')
                        record['formula_status']='saved_result_confirmed'
                        value = cached_cell.value
                        record['value_used']=display_value(value)
                    record['recipe']=parse_bundle(value,**parse_options)
                    record['role']='hole' if record['recipe'] is None else 'sample'
                except ValidationError as exc:
                    record.update(recipe=None,role='invalid',error=str(exc)); errors.append({'cell':cell.coordinate,'error':str(exc)})
                cells.append(record)
        return {'workbook':str(path),'sha256':file_hash(path),'sheet':sheet,
                'range':f'{get_column_letter(minc)}{minr}:{get_column_letter(maxc)}{maxr}',
                'rows':maxr-minr+1,'columns':maxc-minc+1,'cells':cells,'errors':errors,
                'formula_count':formula_count, 'available_sheets':sheets,
                'warnings':(['Formula results are saved by Excel, not calculated by GridCAD. Their freshness cannot be verified; review the displayed saved values.'] if formula_count else []),
                'parser_policy':{'order':parse_options.get('order','CMYKWT'),
                                 'empty_channels_zero':parse_options.get('empty_channels_zero',False),
                                 'sum_tolerance':str(parse_options.get('tolerance',Decimal(0))),
                                 'delimiter':parse_options.get('delimiter','auto'),
                                 'formula_cache_confirmed':formula_cache_confirmed is True}}
    finally:
        wb.close()
        if cached_wb is not None: cached_wb.close()

def transform_position(row, column, rows, columns, transform):
    required={'transpose','mirror_rows','mirror_columns','row_offset','column_offset'}
    if not isinstance(transform,dict) or set(transform) != required:
        raise ValidationError('Explicit transform must contain transpose, mirror_rows, mirror_columns, row_offset, column_offset.')
    for key in ('transpose','mirror_rows','mirror_columns'):
        if not isinstance(transform[key],bool): raise ValidationError(f'{key} must be boolean.')
    for key in ('row_offset','column_offset'):
        if type(transform[key]) is not int: raise ValidationError(f'{key} must be integer.')
    if transform['transpose']:
        row,column=column,row; rows,columns=columns,rows
    if transform['mirror_rows']: row=rows-1-row
    if transform['mirror_columns']: column=columns-1-column
    return row+transform['row_offset'],column+transform['column_offset']

def mapping_preview(grid, project_path, config):
    inv=archive_inventory(project_path); doc=inv['project_document']; p,occ,bodies=project_index(doc)
    if not isinstance(config,dict) or config.get('schema')!='gridcad.mapping.v1':
        raise ValidationError('Mapping schema must be gridcad.mapping.v1.')
    transform=config.get('transform'); transform_position(0,0,grid['rows'],grid['columns'],transform)
    targets=config.get('targets'); _list(targets,'mapping targets')
    issues=list(grid['errors']); warnings=[]; used_ids=set(); samples={}; labels=[]; preserved=[]
    for target in targets:
        if not isinstance(target,dict): raise ValidationError('Target must be an object.')
        role=target.get('role'); ids=target.get('occurrence_ids')
        if role not in ('sample','row_label','column_label','preserved_other'):
            raise ValidationError('Each target needs explicit sample/row_label/column_label/preserved_other role.')
        if not isinstance(ids,list) or not ids or any(not isinstance(x,str) for x in ids):
            raise ValidationError('Target requires nonempty occurrence_ids array.')
        for oid in ids:
            if oid in used_ids: raise ValidationError(f'Occurrence assigned more than once: {oid}')
            if oid not in occ: raise ValidationError(f'Unknown occurrence ID: {oid}')
            used_ids.add(oid)
        if role=='sample':
            if type(target.get('grid_row')) is not int or type(target.get('grid_column')) is not int:
                raise ValidationError('Sample target requires explicit integer grid_row/grid_column.')
            pos=(target['grid_row'],target['grid_column'])
            if pos in samples: raise ValidationError(f'Duplicate physical sample position: {pos}')
            samples[pos]=target
        elif role in ('row_label','column_label'): labels.append(target)
        else: preserved.append(target)
    # A declared leaf/group mapping must not overlap an ancestor and its child.
    for oid in used_ids:
        stack=list(occ[oid].get('children',[]))
        while stack:
            child=stack.pop()
            if child['id'] in used_ids:
                raise ValidationError(f'Mapping overlaps ancestor {oid} and descendant {child["id"]}.')
            stack.extend(child.get('children',[]))
    assignments=[]; matched=set()
    for cell in grid['cells']:
        pos=transform_position(cell['grid_row'],cell['grid_column'],grid['rows'],grid['columns'],transform)
        target=samples.get(pos)
        if cell['role']=='hole':
            if target:
                issues.append({'cell':cell['cell'],'error':f'Blank grid hole maps to a sample target at {pos}; explicit preservation/reclassification needed.'})
            continue
        if cell['role']=='invalid': continue
        if not target:
            issues.append({'cell':cell['cell'],'error':f'No sample target at transformed position {pos}.'}); continue
        matched.add(pos)
        assignments.append({'role':'sample','cell':cell['cell'],'source_grid':[cell['grid_row'],cell['grid_column']],
                            'target_grid':list(pos),'occurrence_ids':target['occurrence_ids'],'recipe':cell['recipe']})
    for pos in set(samples)-matched:
        issues.append({'target_grid':list(pos),'error':'Sample target has no valid occupied source cell.'})
    black=config.get('label_material')
    valid_black=(isinstance(black,dict) and black.get('mode')=='base_material'
                 and str(black.get('resin_id'))=='11936' and black.get('vendor_material_id')=='VEcoFlxBlk' and black.get('name')=='VEcoFlxBlk')
    if labels and not valid_black:
        issues.append({'error':'Labels require explicit Tray Materials EcoFlexBlack identity: base_material, VEcoFlxBlk, 11936. No CMYK approximation allowed.'})
    for target in labels:
        if 'recipe' in target:
            issues.append({'error':'Printable labels use EcoFlexBlack Tray Materials, not a CMYK recipe.', 'occurrence_ids':target['occurrence_ids']})
        assignments.append({'role':target['role'],'occurrence_ids':target['occurrence_ids'],
                            'material_assignment':black if valid_black else None})
    resin_map=config.get('channel_resin_ids',{})
    if not isinstance(resin_map,dict) or set(resin_map)!=set(CHANNELS) or any(not isinstance(x,str) or not x for x in resin_map.values()) or len(set(resin_map.values()))!=6:
        issues.append({'error':'Explicit six distinct channel_resin_ids are missing or ambiguous; tray feasibility is not established.'})
        resin_map={}
    tray_id=config.get('tray_id')
    if not isinstance(tray_id,str) or tray_id not in {t['id'] for t in p['trays']}:
        issues.append({'error':'Select a valid tray_id explicitly.'})
    selected_trays=[t for t in p['trays'] if t['id']==tray_id]
    selected_resins=(selected_trays[0].get('pluginParameters') or {}).get('pj.tray.resins',[]) if len(selected_trays)==1 else []
    if labels and (not isinstance(selected_resins,list) or 'VEcoFlxBlk' not in selected_resins):
        issues.append({'error':'Selected tray does not explicitly contain VEcoFlxBlk. Label assignment is blocked.'})
    for assignment in assignments:
        assignment['current']=[]
        for oid in assignment['occurrence_ids']:
            refs=[(t,b) for t,b in bodies.get(oid,[]) if t['id']==tray_id]
            if len(refs)!=1:
                issues.append({'occurrence_id':oid,'error':f'Expected one actual print body on selected tray, found {len(refs)}.'})
            desired=material_summary(occ[oid].get('material'))
            actual=[material_summary(b.get('material')) for t,b in refs]
            assignment['current'].append({'occurrence_id':oid,'desired':desired,'actual':actual})
            if len(actual)!=1 or desired != actual[0]:
                warnings.append(f'{oid}: desired occurrence and actual print material differ or are absent; recomposition cannot be assumed.')
        assignment['expected_mode']='cmykwt_dm_mixer' if assignment['role']=='sample' else 'base_material'
        assignment['expected_nonzero_resin_ids']=[resin_map[c] for c in CHANNELS if resin_map and Decimal(assignment['recipe']['channels'][c])>0] if assignment['role']=='sample' else ['11936']
    if config.get('mapping_confirmed') is not True:
        issues.append({'error':'Mapping/orientation is not explicitly confirmed. Review the asymmetric preview first.'})
    if labels:
        warnings.append('EcoFlexBlack identity matches observed vendor catalog, but actual selected tray availability is not verified by this adapter. Production remains blocked.')
    for assignment in assignments:
        for current in assignment['current']:
            def matches(m):
                if assignment['role']=='sample':
                    return (m.get('mode')=='cmykwt_dm_mixer' and m.get('recipe') is not None
                        and all(Decimal(m['recipe']['channels'][c])==Decimal(assignment['recipe']['channels'][c]) for c in CHANNELS)
                        and set(m.get('resin_ids',[]))==set(assignment['expected_nonzero_resin_ids']))
                return (m.get('mode')=='base_material' and m.get('vendor_material_id')=='VEcoFlxBlk'
                        and str(m.get('resin_id'))=='11936')
            current['serialized_request_matches']=(matches(current['desired']) and len(current['actual'])==1 and matches(current['actual'][0]))
    unlisted=sorted(set(occ)-used_ids)
    warnings.append('Channel resin IDs are a user-declared mapping, not validation against a live tray inventory or hardware capability.')
    warnings.append('This preview writes no project. Optional unverified candidate copies require a separate schema/profile, geometry-review, material and preservation gate. GUI reopen/composition and print readiness are not established here.')
    return {'status':'PREVIEW_ONLY','application_validated':False,'writer_enabled':False,
            'input_validation_passed':not issues,'physical_mapping_verified':False,'mapping_evidence':'Explicit user/configuration mapping only; no independent geometry matching.','grid':grid,'project_sha256':inv['sha256'],
            'project_version':doc['version'],'transform':transform,'tray_id':tray_id,
            'mapping_confirmed':config.get('mapping_confirmed') is True,
            'assignments':assignments,'preserved_targets':preserved,'unlisted_occurrences_preserved':unlisted,
            'issues':issues,'warnings':warnings,
            'readiness':{'production_write':False,'reason':'Preview success alone is not candidate-copy eligibility or application compatibility; separate admission and preservation checks are required.'}}

def semantic_differences(a,b,path=''):
    # Decimal and int are mathematically equal; strings are never coerced.
    if isinstance(a,dict) and isinstance(b,dict):
        out=[]
        for k in sorted(set(a)|set(b)):
            p=path+'/'+k.replace('~','~0').replace('/','~1')
            if k not in a or k not in b: out.append(p)
            else: out.extend(semantic_differences(a[k],b[k],p))
        return out
    if isinstance(a,list) and isinstance(b,list):
        if len(a)!=len(b): return [path]
        return [p for i,(x,y) in enumerate(zip(a,b)) for p in semantic_differences(x,y,path+'/'+str(i))]
    if type(a) is bool or type(b) is bool:
        return [] if type(a)==type(b) and a==b else [path]
    return [] if a==b else [path]

def audit_projects(before,after):
    a=archive_inventory(before); b=archive_inventory(after)
    ae={x['name']:x for x in a['entries']}; be={x['name']:x for x in b['entries']}
    added=sorted(set(be)-set(ae)); removed=sorted(set(ae)-set(be))
    changed=sorted(k for k in set(ae)&set(be) if k!='project.json' and ae[k]['sha256']!=be[k]['sha256'])
    diffs=semantic_differences(a['project_document'],b['project_document'])
    return {'status':'READ_ONLY_AUDIT','before_sha256':a['sha256'],'after_sha256':b['sha256'],
            'added_archive_entries':added,'removed_archive_entries':removed,'changed_nonproject_payloads':changed,
            'json_changed_paths':diffs,'all_payloads_and_semantics_unchanged':not(added or removed or changed or diffs),
            'application_validated':False,
            'warning':'No changed JSON path is automatically approved. Recipe, geometry, transforms, mode and tray semantics require separate verification.'}

def render_preview(report):
    def e(x): return html.escape(str(x))
    rows=[]; grid=report.get('grid',{}); cells=grid.get('cells',[])
    lookup={(x['grid_row'],x['grid_column']):x for x in cells}
    visual=[]
    if grid.get('rows',0)*grid.get('columns',0)<=2500:
        for r in range(grid.get('rows',0)):
            row=[]
            for c in range(grid.get('columns',0)):
                x=lookup[(r,c)]; pos=transform_position(r,c,grid['rows'],grid['columns'],report['transform'])
                text=x['cell']+' → '+str(pos)+' | '+('HOLE' if x['role']=='hole' else str(x['raw']))
                row.append('<td class="'+e(x['role'])+'">'+e(text)+'</td>')
            visual.append('<tr>'+''.join(row)+'</tr>')
    else: visual=['<tr><td>Large grid: use the full assignment table / JSON instead.</td></tr>']
    for a in report.get('assignments',[]):
        recipe=' / '.join(a['recipe']['channels'][c] for c in CHANNELS) if a['role']=='sample' else 'Tray Materials: EcoFlexBlack (11936), not CMYK'
        matches=sum(x.get('serialized_request_matches',False) for x in a.get('current',[]))
        rows.append('<tr><td>'+e(a['role'])+'</td><td>'+e(a.get('cell','label group'))+'</td><td>'+e(a.get('target_grid',''))+'</td><td>'+e(', '.join(a['occurrence_ids']))+'</td><td>'+e(recipe)+'</td><td>'+e(str(matches)+' / '+str(len(a['occurrence_ids'])))+'</td></tr>')
    return ('<!doctype html><html><head><meta charset="utf-8"><title>GridCAD read-only preview</title><style>body{font:16px system-ui;max-width:1400px;margin:2em auto;padding:1em}table{border-collapse:collapse;width:100%}td,th{border:1px solid #bbb;padding:.5em;text-align:left}.warning{background:#fff1be;padding:1em}.hole{background:#ddd}.invalid{background:#ffd1d1}pre{white-space:pre-wrap}.scroll{overflow:auto}</style></head><body>'
        '<h1>GridCAD read-only preview</h1><p class="warning">No project changed. This is not a GrabCAD-validated production output. Review object identity, labels, orientation, resin identities and all issues. Numerical recipes are NOT displayed as estimated RGB colors.</p>'
        '<h2>Source grid → target coordinates</h2><p>Zero-based target row/column. Empty cells are holes, never zero recipes. This is a mapping diagram, NOT a geometry rendering.</p><div class="scroll"><table>'+''.join(visual)+'</table></div>'
        '<h2>Explicit transform</h2><pre>'+e(json_text(report.get('transform')))+'</pre><h2>Assignments</h2><table><tr><th>Role</th><th>Cell</th><th>Target grid</th><th>Occurrence IDs</th><th>Requested material</th><th>Serialized matches (not app validation)</th></tr>'+''.join(rows)+'</table><h2>Issues</h2><pre>'+e(json_text(report.get('issues',[])))+'</pre><h2>Warnings</h2><pre>'+e(json_text(report.get('warnings',[])))+'</pre><h2>Preserved other targets</h2><pre>'+e(json_text(report.get('preserved_targets',[])))+'</pre></body></html>')
