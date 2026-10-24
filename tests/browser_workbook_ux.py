"""Linux Chromium UX smoke checks. Not Windows, Excel or GrabCAD validation."""
from pathlib import Path
import json
import sys
import tempfile
import threading
from playwright.sync_api import sync_playwright
from grabcad_grid_ui import ReviewServer
from test_workbook_ux import workbook, RECIPE
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
QA = ROOT/'qa'
QA.mkdir(exist_ok=True)
server = ReviewServer()
thread = threading.Thread(target=server.serve_forever, daemon=True)
thread.start()
checks=[]

def check(condition, description):
    assert condition, description
    checks.append(description)

try:
 with tempfile.TemporaryDirectory() as tmp, sync_playwright() as p:
    tmp = Path(tmp)
    cached=tmp/'synthetic-cached.xlsx';workbook(cached,RECIPE)
    missing=tmp/'synthetic-missing.xlsx';workbook(missing)
    longfile=tmp/'synthetic-long.xlsx'
    wb=Workbook();ws=wb.active;ws.title='Long samples'
    for row in range(1,502): ws.cell(row,2,RECIPE if row%7 else None)
    ws['C501']=RECIPE;wb.save(longfile);wb.close()
    browser=p.chromium.launch(executable_path='/usr/bin/chromium',headless=True,args=['--no-sandbox'])
    page=browser.new_page(viewport={'width':1280,'height':900},device_scale_factor=1)
    errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    page.goto(server.origin+'/#'+server.token)
    page.locator('#workbook').set_input_files(str(cached))
    page.wait_for_function("document.querySelector('#sheet-preview button') && !document.querySelector('#analyze').disabled")
    check(page.locator('#sheet').input_value()=='Recipe grid','Worksheet picker uses real tab name')
    check('A1:H3' in page.locator('#sheet-info').inner_text(),'Used area explicitly includes sidebars')
    first=page.locator('#sheet-preview button[data-row="2"][data-column="2"]');last=page.locator('#sheet-preview button[data-row="3"][data-column="3"]')
    first.scroll_into_view_if_needed();a=first.bounding_box();b=last.bounding_box()
    page.mouse.move(a['x']+a['width']/2,a['y']+a['height']/2);page.mouse.down();page.mouse.move(b['x']+b['width']/2,b['y']+b['height']/2,steps=8);page.mouse.up()
    check(page.locator('#range').input_value()=='B2:C3','Pointer drag selects exact addressed rectangle')
    check(page.locator('#sheet-preview button.selected').count()==4,'Selection highlights all cells including blank positions')
    page.locator('#project').set_input_files(str(ROOT/'tests/fixtures/DO_NOT_PRINT_synthetic.PRINT'))
    page.locator('#analyze').click();page.wait_for_selector('#review',state='visible')
    check('freshness cannot be verified' in page.locator('#recipe-errors').inner_text(),'Unconfirmed formula result is blocked with freshness guidance')
    page.locator('#grid button').first.click()
    check('Saved formula result:' in page.locator('#cell-detail').inner_text(),'Recipe detail exposes saved formula value')
    check('Original cell content: =' in page.locator('#cell-detail').inner_text(),'Recipe detail retains formula identity')
    page.locator('#formula-cache').check();page.locator('#analyze').click();page.wait_for_selector('#review',state='visible')
    check(page.locator('#grid button.sample').count()==3,'Confirmed cached result parses alongside literals')
    check('Total: 100%' in page.locator('#cell-detail').inner_text(),'Exact parsed channel sum is visible')
    page.locator('#workbook').set_input_files(str(missing));page.wait_for_function("!document.querySelector('#analyze').disabled && document.querySelector('#sheet-preview button')")
    check(not page.locator('#formula-cache').is_checked(),'Changing workbook resets cache confirmation')
    page.locator('#formula-cache').check();page.locator('#analyze').click();page.wait_for_selector('#review',state='visible')
    check('no readable saved result' in page.locator('#recipe-errors').inner_text(),'Missing cache stays an error, never a hole')
    page.locator('#workbook').set_input_files(str(longfile));page.wait_for_function("document.querySelector('#sheet').value==='Long samples' && !document.querySelector('#analyze').disabled")
    check(page.locator('#sheet-preview button').count()<=1200,'Long-sheet data preview bounds DOM cell count')
    page.locator('#sheet-next').click();page.wait_for_function("document.querySelector('#sheet-preview-status').textContent.includes('A41:')")
    check('A41:' in page.locator('#sheet-preview-status').inner_text(),'Long-sheet next-page navigation retains blank rows')
    page.locator('#preview-start').fill('B481');page.locator('#refresh-sheet').click();page.wait_for_function("document.querySelector('#sheet-preview-status').textContent.includes('B481:')")
    check(page.locator('#sheet-preview button[data-row="501"][data-column="2"]').count()==1,'Direct addressed navigation reaches late worksheet rows')
    page.locator('#range').fill('B1:C501');page.locator('#transpose').check();page.locator('#analyze').click();page.wait_for_selector('#review',state='visible')
    check('source 501 rows × 2 columns' in page.locator('#grid-summary').inner_text(),'Long rectangular source dimensions stay visible')
    check('2 rows × 501 columns (transposed)' in page.locator('#grid-summary').inner_text(),'Explicit transpose swaps preview dimensions without editing source')
    check(page.locator('#grid button').count()==60,'Transposed recipe preview paginates long columns')
    page.locator('#grid-right').click();check('columns 31–60 of 501' in page.locator('#grid-page').inner_text(),'Recipe preview pages across transposed long dimension')
    # Complete unchanged synthetic mapping and exact candidate preview.
    page.locator('#workbook').set_input_files(str(ROOT/'tests/fixtures/DO_NOT_PRINT_grid.xlsx'));page.wait_for_function("document.querySelector('#sheet').value==='Samples' && !document.querySelector('#analyze').disabled")
    page.locator('#range').fill('B2:E3');page.locator('#transpose').uncheck();page.locator('#analyze').click();page.wait_for_selector('#review',state='visible')
    page.locator('#region-sample').fill('29,39,83,67')
    # Input changes intentionally hide stale review. Use main Inspect button.
    page.locator('#analyze').click();page.wait_for_selector('#review',state='visible')
    page.locator('#region-row_label').fill('11,39,17,67');page.locator('#analyze').click();page.wait_for_selector('#review',state='visible')
    page.locator('#region-column_label').fill('29,27,86,30');page.locator('#analyze').click();page.wait_for_selector('#review',state='visible')
    check(page.locator('#candidate').input_value()=='','Orientation is never silently selected')
    options=page.locator('#candidate option').evaluate_all('(els)=>els.map(e=>({value:e.value,label:e.textContent}))')
    chosen=next(o for o in options if 'Original axes' in o['label'])
    page.locator('#candidate').select_option(chosen['value']);page.locator('#reviewed').check();page.locator('#overwrite').check();page.locator('#compare').click();page.wait_for_function("!document.querySelector('#save').disabled")
    check('bodies will change' in page.locator('#comparison').inner_text(),'Reviewed exact-change preview enables separate-copy save')
    page.locator('#sheet-selector').scroll_into_view_if_needed();page.screenshot(path=str(QA/'worksheet-selector.png'),full_page=True)
    page.locator('#candidate').scroll_into_view_if_needed();page.screenshot(path=str(QA/'mapping-preview.png'))
    check(not errors,'No uncaught browser JavaScript errors')
    browser.close()
finally:
 server.shutdown();server.server_close();thread.join()
print(json.dumps({'platform':'Linux Chromium','checks_passed':len(checks),'checks':checks},indent=2))
