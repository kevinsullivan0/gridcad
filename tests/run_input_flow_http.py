"""Run production JS + HTTP integration in one permitted executor namespace."""
from pathlib import Path
import tempfile,threading,subprocess,json,hashlib
from openpyxl import load_workbook, Workbook
from grabcad_grid_ui import ReviewServer
from gridcad.core import audit_projects,archive_inventory,project_index,semantic_differences
from test_workbook_ux import workbook,RECIPE
from test_input_flow_api import save_workbook, inject_cache
ROOT=Path(__file__).resolve().parents[1]
with tempfile.TemporaryDirectory() as temp:
    p=Path(temp)
    wb=load_workbook(ROOT/'tests/fixtures/DO_NOT_PRINT_grid.xlsx')
    ws=wb.active;ws.title='Recipes';ws['A2']=20
    settings=wb.create_sheet('Settings',0);settings['A1']='Sidebar only'
    long=wb.create_sheet('Long')
    for row in range(2,502):
        if row%7:long.cell(row,2,RECIPE)
    long['C501']=RECIPE
    wb.active=1;wb.save(p/'multi.xlsx');wb.close()
    workbook(p/'cached.xlsx',RECIPE)
    empty=Workbook();empty.save(p/'empty.xlsx');empty.close()
    save_workbook(p/'bad-cache.xlsx',formula=True);inject_cache(p/'bad-cache.xlsx','999',cache_type='s')
    save_workbook(p/'infinite-cache.xlsx',formula=True);inject_cache(p/'infinite-cache.xlsx','1e999',cache_type='n')
    (p/'synthetic.PRINT').write_bytes((ROOT/'tests/fixtures/DO_NOT_PRINT_synthetic.PRINT').read_bytes())
    original=(p/'synthetic.PRINT').read_bytes()
    server=ReviewServer();thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    try:
        result=subprocess.run(['node',str(ROOT/'tests/test_input_flow_http.js'),server.origin,server.token,str(p)],text=True,capture_output=True,timeout=120)
        if result.returncode:
            print(result.stdout,end='');print(result.stderr,end='');raise SystemExit(result.returncode)
        report=json.loads(result.stdout)
        downloaded=p/'downloaded.PRINT'
        assert downloaded.is_file()
        audit=audit_projects(p/'synthetic.PRINT',downloaded)
        assert not any(audit[k] for k in ('added_archive_entries','removed_archive_entries','changed_nonproject_payloads'))
        assert (p/'synthetic.PRINT').read_bytes()==original
        before=archive_inventory(p/'synthetic.PRINT')['project_document'];after=archive_inventory(downloaded)['project_document']
        for doc in (before,after):
            _,nodes,bodies=project_index(doc)
            for node in nodes.values():node.pop('material',None)
            for refs in bodies.values():
                for tray,body in refs:body.pop('material',None)
        assert semantic_differences(before,after)==[]
        report['download_audit']={'original_unchanged':True,'archive_payloads_preserved':True,'nonmaterial_json_preserved':True,'download_bytes':downloaded.stat().st_size,'download_sha256':hashlib.sha256(downloaded.read_bytes()).hexdigest()}
        print(json.dumps(report,indent=2))
    finally:server.shutdown();server.server_close();thread.join()
