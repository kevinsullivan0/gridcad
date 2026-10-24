"""New workbook UX regression checks, not the unrecovered release suite.

All XLSX files are synthetic and temporary. Formula caches are deliberately
constructed XML fixtures, not evidence that any Excel engine evaluated them.
"""
import json
from pathlib import Path
import tempfile
import threading
import unittest
from urllib.request import Request, urlopen
from urllib.error import HTTPError
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook
from gridcad.core import ValidationError, read_workbook, parse_bundle, file_hash
from grabcad_grid_ui import ReviewServer

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
RECIPE = '20 / 0 / 0 / 0 / 30 / 50'


def workbook(path, cache=None, cache_type='str', second_sheet=False):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Recipe grid'
    ws['A1'] = 'Sidebar / heading, not a recipe'
    ws['A2'] = 20
    ws['B2'] = '=A2&" / 0 / 0 / 0 / 30 / 50"'
    ws['C2'] = '  0 / 10 / 0 / 0 / 30 / 60  '
    ws['C3'] = '0;0;10;0;30;60'
    ws['H2'] = 'Sidebar input'
    if second_sheet:
        wb.create_sheet('Settings')
    wb.save(path)
    wb.close()
    if cache is not None:
        with ZipFile(path) as z:
            entries = [(i, z.read(i.filename)) for i in z.infolist()]
        with ZipFile(path, 'w') as z:
            for info, data in entries:
                if info.filename == 'xl/worksheets/sheet1.xml':
                    root = ET.fromstring(data)
                    cell = next(c for c in root.iter(NS+'c') if c.attrib.get('r') == 'B2')
                    cell.attrib['t'] = cache_type
                    v = cell.find(NS+'v')
                    if v is None: v = ET.SubElement(cell, NS+'v')
                    v.text = cache
                    data = ET.tostring(root, encoding='utf-8')
                z.writestr(info, data)


class WorkbookUXTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / 'Example with spaces.xlsx'
    def tearDown(self): self.tmp.cleanup()
    def read(self, **kw): return read_workbook(self.path, sheet='Recipe grid', cell_range='B2:C3', **kw)
    def test_cached_formula_requires_explicit_confirmation(self):
        workbook(self.path, RECIPE)
        c = self.read()['cells'][0]
        self.assertEqual(c['role'], 'invalid')
        self.assertEqual(c['cached_value'], RECIPE)
        self.assertEqual(c['formula_status'], 'confirmation_required')
        self.assertIsNone(c['value_used'])
    def test_confirmed_cached_result_preserves_formula_and_exact_channels(self):
        workbook(self.path, RECIPE)
        before = self.path.read_bytes()
        g = self.read(formula_cache_confirmed=True)
        c = g['cells'][0]
        self.assertEqual(c['role'], 'sample')
        self.assertTrue(c['raw'].startswith('='))
        self.assertEqual(c['value_used'], RECIPE)
        self.assertEqual(c['recipe']['channels'], dict(zip('CMYKWT', ['20','0','0','0','30','50'])))
        self.assertEqual(c['recipe']['total'], '100')
        self.assertEqual(c['cache_freshness'], 'not_verifiable')
        self.assertEqual(g['formula_count'], 1)
        self.assertTrue(g['warnings'])
        self.assertEqual(g['cells'][2]['role'], 'hole')
        self.assertFalse(g['errors'])
        self.assertEqual(before, self.path.read_bytes())
    def test_missing_cache_is_invalid_not_a_hole(self):
        workbook(self.path)
        for confirmed in (False, True):
            c = self.read(formula_cache_confirmed=confirmed)['cells'][0]
            self.assertEqual(c['role'], 'invalid')
            self.assertEqual(c['formula_status'], 'missing_saved_result')
            self.assertIn('recalculate', c['error'])
            self.assertIn('values in a workbook copy', c['error'])
    def test_formula_excel_error_reports_specific_error(self):
        workbook(self.path, '#VALUE!', 'e')
        c = self.read(formula_cache_confirmed=True)['cells'][0]
        self.assertEqual(c['formula_status'], 'saved_excel_error')
        self.assertIn('#VALUE!', c['error'])
    def test_numeric_formula_result_is_not_a_bundle(self):
        workbook(self.path, '100', 'n')
        c = self.read(formula_cache_confirmed=True)['cells'][0]
        self.assertEqual(c['role'], 'invalid')
        self.assertIn('not a scalar', c['error'])
    def test_bad_cached_format_is_distinct_from_missing_cache(self):
        workbook(self.path, '20/30/50')
        c = self.read(formula_cache_confirmed=True)['cells'][0]
        self.assertEqual(c['formula_status'], 'saved_result_confirmed')
        self.assertIn('six positions', c['error'])
        self.assertNotIn('no readable saved result', c['error'])
    def test_cache_confirmation_must_be_true_boolean(self):
        workbook(self.path, RECIPE)
        for fake in ('true', 1, None):
            self.assertEqual(self.read(formula_cache_confirmed=fake)['cells'][0]['role'], 'invalid')
    def test_metadata_lists_tabs_without_recipe_inference(self):
        workbook(self.path, second_sheet=True)
        meta = read_workbook(self.path, metadata_only=True)
        self.assertEqual(meta['sheets'], [{'name':'Recipe grid','used_range':'A1:H3'}, {'name':'Settings','used_range':'A1:A1'}])
        self.assertFalse(meta['source_modified'])
        with self.assertRaisesRegex(ValidationError, 'Available tabs'):
            read_workbook(self.path, cell_range='B2:C3')
    def test_sidebars_excluded_and_blank_positions_retained(self):
        workbook(self.path, RECIPE)
        g = self.read(formula_cache_confirmed=True)
        self.assertEqual((g['rows'], g['columns']), (2,2))
        self.assertEqual([c['cell'] for c in g['cells']], ['B2','C2','B3','C3'])
        self.assertEqual(g['cells'][2]['grid_row'],1)
        self.assertEqual(g['cells'][2]['grid_column'],0)
        self.assertEqual(g['cells'][2]['role'],'hole')
        broad = read_workbook(self.path, sheet='Recipe grid', cell_range='A1:H3', formula_cache_confirmed=True)
        self.assertTrue(any(e['cell']=='A1' for e in broad['errors']))
    def test_array_spill_anchor_outside_selection_is_rejected(self):
        from openpyxl.worksheet.formula import ArrayFormula
        wb=Workbook();ws=wb.active;ws.title='Recipe grid'
        ws['A1']=ArrayFormula(ref='A1:B1', text='=IF({1,2},"20/0/0/0/30/50")')
        ws['B1']=RECIPE
        wb.save(self.path);wb.close()
        for confirmed in (False, True):
            with self.assertRaisesRegex(ValidationError, 'array/spill'):
                read_workbook(self.path,sheet='Recipe grid',cell_range='B1',formula_cache_confirmed=confirmed)
        preview=read_workbook(self.path,sheet='Recipe grid',cell_range='B1',metadata_only=True)
        self.assertEqual(preview['preview_cells'][0]['unsupported_formula_region'], 'A1 (A1:B1)')
    def test_array_outside_recipe_selection_does_not_block_literals(self):
        from openpyxl.worksheet.formula import ArrayFormula
        wb=Workbook();ws=wb.active;ws.title='Recipe grid'
        ws['A1']=ArrayFormula(ref='A1:B1',text='=IF({1,2},"20/0/0/0/30/50")')
        ws['B2']=RECIPE;wb.save(self.path);wb.close()
        grid=read_workbook(self.path,sheet='Recipe grid',cell_range='B2')
        self.assertEqual(grid['cells'][0]['role'],'sample')
    def test_data_table_anchor_outside_selection_is_rejected(self):
        from openpyxl.worksheet.formula import DataTableFormula
        wb=Workbook();ws=wb.active;ws.title='Recipe grid'
        ws['A1']=DataTableFormula(ref='A1:B1', r1='D1')
        ws['B1']=RECIPE;wb.save(self.path);wb.close()
        for confirmed in (False, True):
            with self.assertRaisesRegex(ValidationError, 'data-table'):
                read_workbook(self.path,sheet='Recipe grid',cell_range='B1',formula_cache_confirmed=confirmed)
        preview=read_workbook(self.path,sheet='Recipe grid',cell_range='B1',metadata_only=True)
        self.assertEqual(preview['preview_cells'][0]['unsupported_formula_region'], 'A1 (A1:B1)')
    def test_large_data_preview_is_bounded_and_retains_addresses(self):
        wb=Workbook();ws=wb.active;ws.title='Recipe grid';ws['AZ700']=RECIPE;wb.save(self.path);wb.close()
        meta=read_workbook(self.path, metadata_only=True, sheet='Recipe grid',cell_range='Z650')
        self.assertLessEqual(len(meta['preview_cells']),1200)
        self.assertEqual(meta['preview_cells'][0]['cell'],'Z650')
        self.assertTrue(meta['preview_limited'])
    def test_preview_rejects_invalid_origin(self):
        workbook(self.path)
        for origin in ('A0','A:A','A1:B2','XFE1'):
            with self.subTest(origin=origin), self.assertRaises(ValidationError):
                read_workbook(self.path,metadata_only=True,sheet='Recipe grid',cell_range=origin)
    def test_spaces_slashes_decimals_and_percent_signs(self):
        for value in (' 20 / 0 / 0 / 0 / 30 / 50 ', '20 0 0 0 30 50', '20;0;0;0;30;50', '(20% / 0% / 0% / 0% / 30% / 50%)'):
            self.assertEqual(parse_bundle(value)['total'],'100')
        self.assertEqual(parse_bundle('0.0001/0/0/0/0/99.9999')['channels']['C'],'0.0001')
    def test_ambiguity_and_implicit_zero_still_rejected(self):
        for value in ('20/0;0/0/30/50', '20,0,0,0,30,50', '20///0/30/50', '20/0/0/30/50', '20/0/0/0/30/49.9999'):
            with self.subTest(value=value), self.assertRaises(ValidationError): parse_bundle(value)
        self.assertEqual(parse_bundle('20///0/30/50',empty_channels_zero=True)['total'],'100')


class WorkbookAPITests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name)/'synthetic.xlsx'
        workbook(self.path, RECIPE, second_sheet=True)
        self.server = ReviewServer()
        self.worker = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.worker.start()
    def tearDown(self):
        self.server.shutdown(); self.server.server_close(); self.worker.join(); self.tmp.cleanup()
    def post(self, path, data, binary=False):
        req = Request(self.server.origin+path, data=data if binary else json.dumps(data).encode(), headers={'Content-Type':'application/octet-stream' if binary else 'application/json','X-GridCAD-Token':self.server.token}, method='POST')
        try:
            with urlopen(req) as res: return res.status, json.load(res)
        except HTTPError as exc: return exc.code, json.load(exc)
    def test_workbook_upload_returns_real_tabs(self):
        status, value = self.post('/api/upload/workbook',self.path.read_bytes(), True)
        self.assertEqual(status,200)
        self.assertEqual(value['workbook_info']['sheets'][0]['name'],'Recipe grid')
        self.assertFalse(value['source_modified'])
    def test_api_requires_deliberate_rectangle(self):
        self.post('/api/upload/workbook',self.path.read_bytes(),True)
        self.post('/api/upload/project',b'placeholder',True)
        status, value = self.post('/api/analyze',{'sheet':'Recipe grid'})
        self.assertEqual(status,400)
        self.assertIn('Enter the recipe grid rectangle',value['error'])
    def test_unconfirmed_cache_diagnostics_do_not_enable_mapping(self):
        fixture=Path(__file__).parent/'fixtures/DO_NOT_PRINT_synthetic.PRINT'
        self.post('/api/upload/workbook',self.path.read_bytes(),True)
        self.post('/api/upload/project',fixture.read_bytes(),True)
        status, value = self.post('/api/analyze',{'sheet':'Recipe grid','range':'B2:C3'})
        self.assertEqual(status,200)
        self.assertEqual(value['grid']['cells'][0]['formula_status'],'confirmation_required')
        self.assertFalse((value.get('proposal') or {}).get('orientation_candidates'))

if __name__=='__main__': unittest.main()
