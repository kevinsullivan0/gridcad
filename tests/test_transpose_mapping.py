"""NEW transpose/long-sheet tests using synthetic inputs only.

These extend the newly written regression checks; they are not recovered
release tests. The fixture is the hash-verified DO_NOT_PRINT release example.
No printable output or text/label inference is established by these tests.
"""
import copy
import json
from pathlib import Path
import tempfile
import threading
import unittest
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from openpyxl import Workbook

from grabcad_grid_ui import ReviewServer
from gridcad import core
from gridcad.core import (
    ValidationError, archive_inventory, audit_projects, file_hash, project_index,
    read_workbook, semantic_differences,
)
from gridcad.geometry import propose_grid_mapping
from gridcad.candidate_writer import mapping_from_geometry_review, write_reviewed_candidate_copy
from gridcad.material_converter import EcoFlexMaterialProvider


FIXTURES = Path(__file__).parent / "fixtures"
SOURCE = FIXTURES / "DO_NOT_PRINT_synthetic.PRINT"
SOURCE_WORKBOOK = FIXTURES / "DO_NOT_PRINT_grid.xlsx"
SELECTION = {
    "region_mm": [29, 39, 83, 67],
    "role_regions_mm": {
        "row_label": [11, 39, 17, 67],
        "column_label": [29, 27, 86, 30],
    },
}
TRANSPOSE = dict(transpose=True, mirror_rows=False, mirror_columns=False,
                 row_offset=0, column_offset=0)
EXPECTED = {
    "C5": (0, 0), "D5": (1, 0),
    "C6": (0, 1), "D6": (1, 1),
    "C8": (0, 3), "D8": (1, 3),
    "D7": (1, 2),
}


def transposed_workbook(path):
    """Transpose source 2x4 values into C5:D8, preserving one blank at C7."""
    original = read_workbook(SOURCE_WORKBOOK, sheet="Samples", cell_range="B2:E3")
    wb = Workbook()
    ws = wb.active
    ws.title = "Long-axis recipes"
    ws["A1"] = "Synthetic sidebar, not a recipe"
    for cell in original["cells"]:
        ws.cell(row=5 + cell["grid_column"], column=3 + cell["grid_row"], value=cell["raw"])
    wb.save(path)
    wb.close()


def long_workbook(path, invalid=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Long recipes"
    ws["A1"] = "Synthetic sidebar"
    ws["B3"] = "0.125 / 0 / 0 / 0 / 0 / 99.875"
    ws["E502"] = "0 / 25 / 25 / 0 / 50 / 0"
    if invalid:
        ws["E501"] = 100
    wb.save(path)
    wb.close()


class TransposedFixture:
    def make_transposed(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.workbook = self.root / "synthetic_transposed.xlsx"
        transposed_workbook(self.workbook)
        self.grid = read_workbook(self.workbook, sheet="Long-axis recipes", cell_range="C5:D8")
        self.proposal = propose_grid_mapping(self.grid, SOURCE, selection=copy.deepcopy(SELECTION))
        self.candidate = next(c for c in self.proposal["orientation_candidates"] if c["transform"] == TRANSPOSE)


class TransposeMappingTests(TransposedFixture, unittest.TestCase):
    def setUp(self):
        self.make_transposed()

    def test_original_4_by_2_cell_addresses_map_to_physical_2_by_4(self):
        self.assertEqual((self.grid["rows"], self.grid["columns"]), (4, 2))
        self.assertEqual(self.grid["range"], "C5:D8")
        self.assertEqual([c["cell"] for c in self.grid["cells"]],
                         ["C5", "D5", "C6", "D6", "C7", "D7", "C8", "D8"])
        self.assertEqual(len(self.proposal["orientation_candidates"]), 1)
        self.assertEqual(self.candidate["transform"], TRANSPOSE)
        cells = {c["cell"]: c for c in self.grid["cells"]}
        self.assertEqual({a["cell"]: tuple(a["target_grid"]) for a in self.candidate["assignments"]}, EXPECTED)
        for assignment in self.candidate["assignments"]:
            cell = cells[assignment["cell"]]
            self.assertEqual(assignment["source_grid"], [cell["grid_row"], cell["grid_column"]])
            r, c = EXPECTED[assignment["cell"]]
            self.assertEqual(assignment["occurrence_id"], f"case_profile-sample-r{r}-c{c}")

    def test_hole_moves_without_becoming_a_sample_or_losing_its_address(self):
        hole = next(c for c in self.grid["cells"] if c["role"] == "hole")
        self.assertEqual((hole["cell"], hole["grid_row"], hole["grid_column"]), ("C7", 2, 0))
        self.assertIsNone(hole["recipe"])
        self.assertEqual(self.candidate["holes"],
                         [{"source_grid": [2, 0], "target_grid": [0, 2], "occurrence_ids": []}])
        self.assertNotIn("C7", {a["cell"] for a in self.candidate["assignments"]})
        self.assertNotIn((0, 2), {tuple(a["target_grid"]) for a in self.candidate["assignments"]})

    def test_physical_label_regions_and_direct_black_roles_do_not_swap(self):
        normal = read_workbook(SOURCE_WORKBOOK, sheet="Samples", cell_range="B2:E3")
        normal_proposal = propose_grid_mapping(normal, SOURCE, selection=copy.deepcopy(SELECTION))
        normal_targets = normal_proposal["orientation_candidates"][0]["targets"]
        normal_labels = {t["role"]: t for t in normal_targets if t["role"] != "sample"}
        transposed_labels = {t["role"]: t for t in self.candidate["targets"] if t["role"] != "sample"}
        self.assertEqual(transposed_labels, normal_labels)
        self.assertEqual(self.proposal["selected_label_counts"], {"row_label": 4, "column_label": 8})
        self.assertEqual(self.candidate["label_text_association"], "not_inferred")
        for role, target in transposed_labels.items():
            self.assertTrue(all(oid.startswith(role + "-") for oid in target["occurrence_ids"]))
            self.assertEqual(target["material"], {"mode": "base_material", "name": "VEcoFlxBlk",
                                                "vendor_material_id": "VEcoFlxBlk", "resin_id": "11936"})
            self.assertIsNone(target["label_text"])
            self.assertEqual(target["text_association"], "not_inferred")
            self.assertNotIn("recipe", target)

    def test_transpose_does_not_silently_confirm_orientation_or_labels(self):
        self.assertFalse(self.proposal["mapping_confirmed"])
        self.assertFalse(self.proposal["writer_enabled"])
        self.assertFalse(self.candidate["mapping_confirmed"])
        self.assertFalse(self.candidate["label_grouping_confirmed"])
        with self.assertRaises(ValidationError):
            mapping_from_geometry_review(self.grid, SOURCE, selection=copy.deepcopy(SELECTION),
                                         candidate_id=self.candidate["candidate_id"], reviewed=False)

    def test_transposed_candidate_materials_and_every_nonmaterial_field_preserved(self):
        before_source = file_hash(SOURCE)
        before_workbook = self.workbook.read_bytes()
        source_grid = copy.deepcopy(self.grid)
        output = self.root / "DO_NOT_PRINT_unverified_transposed.PRINT"
        result = write_reviewed_candidate_copy(self.grid, SOURCE, output,
            selection=copy.deepcopy(SELECTION), candidate_id=self.candidate["candidate_id"],
            assignment_policy={"mode": "overwrite_targets"}, reviewed=True)
        self.assertEqual(file_hash(SOURCE), before_source)
        self.assertEqual(self.workbook.read_bytes(), before_workbook)
        self.assertEqual(self.grid, source_grid)
        self.assertEqual(result["status"], "UNVERIFIED_CANDIDATE_COPY")
        self.assertFalse(result["print_ready"])
        self.assertFalse(result["production_write"])
        self.assertFalse(result["gui_validated"])
        self.assertTrue(result["geometry_mapping_recomputed_passed"])
        audit = audit_projects(SOURCE, output)
        for key in ("added_archive_entries", "removed_archive_entries", "changed_nonproject_payloads"):
            self.assertEqual(audit[key], [])
        after = archive_inventory(output)["project_document"]
        _, nodes, bodies = project_index(after)
        provider = EcoFlexMaterialProvider()
        cells = {c["cell"]: c for c in self.grid["cells"]}
        for assignment in self.candidate["assignments"]:
            oid = assignment["occurrence_id"]
            expected = provider.sample(cells[assignment["cell"]]["recipe"])
            for actual in (nodes[oid]["material"], bodies[oid][0][1]["material"]):
                self.assertEqual({k: v for k, v in actual.items() if k != "id"},
                                 {k: v for k, v in expected.items() if k != "id"})
        for target in self.candidate["targets"]:
            if target["role"] == "sample":
                continue
            for oid in target["occurrence_ids"]:
                for material in (nodes[oid]["material"], bodies[oid][0][1]["material"]):
                    self.assertEqual(material["metadata"]["materialType"], "base_material")
                    self.assertEqual(material["metadata"]["resinId"], 11936)
                    self.assertEqual(material["vendorMaterialId"], "VEcoFlxBlk")
        before = archive_inventory(SOURCE)["project_document"]
        for doc in (before, after):
            _, nodes, bodies = project_index(doc)
            for oid in result["changed_occurrence_ids"]:
                nodes[oid].pop("material")
                bodies[oid][0][1].pop("material")
        self.assertEqual(semantic_differences(before, after), [])


class LongWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / "synthetic_500x4.xlsx"

    def test_500_by_4_selection_keeps_all_empty_rows_and_columns(self):
        long_workbook(self.path)
        before = self.path.read_bytes()
        grid = read_workbook(self.path, sheet="Long recipes", cell_range="B3:E502")
        self.assertEqual((grid["rows"], grid["columns"]), (500, 4))
        self.assertEqual(len(grid["cells"]), 2000)
        self.assertEqual(grid["errors"], [])
        self.assertEqual(sum(c["role"] == "hole" for c in grid["cells"]), 1998)
        samples = [c for c in grid["cells"] if c["role"] == "sample"]
        self.assertEqual([(c["cell"], c["grid_row"], c["grid_column"]) for c in samples],
                         [("B3", 0, 0), ("E502", 499, 3)])
        self.assertEqual(samples[0]["recipe"]["channels"]["C"], "0.125")
        self.assertEqual(samples[0]["recipe"]["channels"]["T"], "99.875")
        for column in (1, 2):
            cells = [c for c in grid["cells"] if c["grid_column"] == column]
            self.assertEqual(len(cells), 500)
            self.assertTrue(all(c["role"] == "hole" for c in cells))
        middle = [c for c in grid["cells"] if c["grid_row"] == 250]
        self.assertEqual(len(middle), 4)
        self.assertTrue(all(c["role"] == "hole" for c in middle))
        self.assertEqual(self.path.read_bytes(), before)

    def test_validation_reaches_last_page_and_does_not_turn_errors_into_holes(self):
        long_workbook(self.path, invalid=True)
        grid = read_workbook(self.path, sheet="Long recipes", cell_range="B3:E502")
        self.assertEqual(len(grid["cells"]), 2000)
        self.assertEqual([error["cell"] for error in grid["errors"]], ["E501"])
        invalid = next(c for c in grid["cells"] if c["cell"] == "E501")
        self.assertEqual((invalid["grid_row"], invalid["grid_column"], invalid["role"]), (498, 3, "invalid"))
        self.assertIn("not a scalar", invalid["error"])
        self.assertEqual(grid["cells"][-1]["cell"], "E502")
        self.assertEqual(grid["cells"][-1]["role"], "sample")

    def test_long_sheet_preview_window_keeps_original_end_page_addresses(self):
        long_workbook(self.path)
        page = read_workbook(self.path, metadata_only=True, sheet="Long recipes", cell_range="B463")
        self.assertEqual(page["preview_range"], "B463:E502")
        self.assertEqual((page["rows"], page["columns"]), (40, 4))
        self.assertEqual(len(page["preview_cells"]), 160)
        self.assertEqual(page["preview_cells"][0]["cell"], "B463")
        self.assertEqual(page["preview_cells"][-1]["cell"], "E502")
        self.assertEqual(page["preview_cells"][-1]["raw"], "0 / 25 / 25 / 0 / 50 / 0")
        self.assertTrue(page["preview_limited"])
        self.assertFalse(page["source_modified"])

    def test_selected_rectangles_above_100000_cells_are_rejected(self):
        long_workbook(self.path)
        self.assertEqual(core.MAX_CELLS, 100000)
        for selected in ("A1:A100001", "A1:GR501"):
            with self.subTest(selected=selected), self.assertRaisesRegex(ValidationError, "Selected range exceeds cell limit"):
                read_workbook(self.path, sheet="Long recipes", cell_range=selected)


class TransposeAPITests(TransposedFixture, unittest.TestCase):
    def setUp(self):
        self.make_transposed()
        self.server = ReviewServer()
        self.worker = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.worker.start()
        self.addCleanup(self.close_server)
        self.assertEqual(self.post("/api/upload/workbook", self.workbook.read_bytes(), binary=True)[0], 200)
        self.assertEqual(self.post("/api/upload/project", SOURCE.read_bytes(), binary=True)[0], 200)

    def close_server(self):
        self.server.shutdown()
        self.server.server_close()
        self.worker.join()

    def post(self, path, data, binary=False):
        request = Request(self.server.origin + path, data=data if binary else json.dumps(data).encode(),
                          headers={"Content-Type": "application/octet-stream" if binary else "application/json",
                                   "X-GridCAD-Token": self.server.token}, method="POST")
        try:
            with urlopen(request) as response:
                return response.status, json.load(response)
        except HTTPError as error:
            return error.code, json.load(error)

    def analyze(self, transpose):
        return self.post("/api/analyze", {"sheet": "Long-axis recipes", "range": "C5:D8",
            "selection": copy.deepcopy(SELECTION), "transpose": transpose})

    def test_explicit_toggle_filters_candidates_without_transposing_source_data(self):
        status, off = self.analyze(False)
        self.assertEqual(status, 200)
        self.assertEqual(off["proposal"]["orientation_candidates"], [])
        self.assertTrue(any("transpose setting" in issue for issue in off["proposal"]["issues"]))
        status, on = self.analyze(True)
        self.assertEqual(status, 200)
        self.assertTrue(on["transpose"])
        self.assertEqual((on["grid"]["rows"], on["grid"]["columns"]), (4, 2))
        self.assertEqual(on["grid"]["cells"], off["grid"]["cells"])
        candidates = on["proposal"]["orientation_candidates"]
        self.assertEqual(len(candidates), 1)
        self.assertEqual(candidates[0]["transform"], TRANSPOSE)
        self.assertEqual({a["cell"]: tuple(a["target_grid"]) for a in candidates[0]["assignments"]}, EXPECTED)
        self.assertFalse(on["proposal"]["mapping_confirmed"])
        self.assertIsNone(self.server.review)

    def test_transpose_accepts_only_explicit_boolean(self):
        for transpose in ("true", 1, None):
            status, result = self.analyze(transpose)
            self.assertEqual(status, 400)
            self.assertIn("explicit on/off", result["error"])
            self.assertIsNone(self.server.review)


if __name__ == "__main__":
    unittest.main()
