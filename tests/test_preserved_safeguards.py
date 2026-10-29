"""NEW regression coverage, written 2026-10-24; not the missing original suite.

All inputs are generated here or are the shipped, explicitly synthetic
DO_NOT_PRINT demo. No private workbook, vendor runtime, or printer is used.
See REGRESSION_PROVENANCE.md for source hashes and coverage boundaries.
"""
import copy
import builtins
import itertools
import os
from pathlib import Path
import stat
import tempfile
import unittest
from unittest.mock import patch
import warnings
import zipfile
from decimal import Decimal

from openpyxl import Workbook

from gridcad import core
from gridcad.core import (
    ValidationError, archive_inventory, audit_projects, file_hash, mapping_preview,
    parse_bundle, project_index, read_workbook, semantic_differences, strict_json,
    transform_position, validate_recipe,
)
from gridcad.fixture_writer import exact_json
from gridcad.geometry import inspect_geometry, propose_grid_mapping, read_stl
from gridcad.material_converter import CHANNEL_RESIN_IDS, EcoFlexMaterialProvider
from gridcad.candidate_writer import (
    candidate_readiness, mapping_from_geometry_review, write_candidate_copy,
    write_reviewed_candidate_copy,
)


FIXTURES = Path(__file__).parent / "fixtures"
SOURCE = FIXTURES / "DO_NOT_PRINT_synthetic.PRINT"
WORKBOOK = FIXTURES / "DO_NOT_PRINT_grid.xlsx"
SELECTION = {
    "region_mm": [29, 39, 83, 67],
    "role_regions_mm": {
        "row_label": [11, 39, 17, 67],
        "column_label": [29, 27, 86, 30],
    },
}
IDENTITY = dict(transpose=False, mirror_rows=False, mirror_columns=False,
                row_offset=0, column_offset=0)


class TempTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def archive(self, entries, name="test.PRINT"):
        path = self.root / name
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            with zipfile.ZipFile(path, "w") as z:
                for key, value in entries:
                    if isinstance(key,str):
                        # ZipInfo sanitizes backslashes on Windows at creation.
                        # These refusal fixtures must contain the requested raw
                        # member name, not an already repaired safe path.
                        info=zipfile.ZipInfo(key)
                        info.filename=key;info.orig_filename=key
                        key=info
                    z.writestr(key, value)
        return path

    def mutate_project(self, mutate, name="mutated.PRINT"):
        with zipfile.ZipFile(SOURCE) as z:
            entries = [(info, z.read(info)) for info in z.infolist()]
        doc = strict_json(next(data for info, data in entries if info.filename == "project.json"))
        mutate(doc)
        return self.archive([(info, exact_json(doc) if info.filename == "project.json" else data)
                             for info, data in entries], name)


class RecipeSafeguards(unittest.TestCase):
    def test_blank_is_hole_not_zero_recipe(self):
        for value in (None, "", " \t "):
            self.assertIsNone(parse_bundle(value))
        with self.assertRaises(ValidationError):
            parse_bundle("0/0/0/0/0/0")

    def test_decimal_precision_is_preserved(self):
        recipe = parse_bundle("0.1234567890123456789/0/0/0/0/99.8765432109876543211")
        self.assertEqual(recipe["channels"]["C"], "0.1234567890123456789")
        self.assertEqual(Decimal(recipe["total"]), Decimal(100))

    def test_tolerance_does_not_normalize(self):
        recipe = parse_bundle("99.99/0/0/0/0/0", tolerance=Decimal(".01"))
        self.assertEqual(recipe["channels"]["C"], "99.99")
        self.assertEqual(recipe["total"], "99.99")
        with self.assertRaises(ValidationError):
            parse_bundle("99.99/0/0/0/0/0")

    def test_channel_order_is_explicit(self):
        self.assertEqual(parse_bundle("100/0/0/0/0/0", order="TCMYKW")["channels"]["T"], "100")
        with self.assertRaises(ValidationError):
            parse_bundle("100/0/0/0/0/0", order="CCMYKW")

    def test_empty_channel_needs_explicit_opt_in(self):
        with self.assertRaises(ValidationError):
            parse_bundle("100/////")
        self.assertEqual(parse_bundle("100/////", empty_channels_zero=True)["channels"]["T"], "0")
        with self.assertRaises(ValidationError):
            parse_bundle("100////", empty_channels_zero=True)

    def test_comma_is_not_silently_inferred(self):
        with self.assertRaises(ValidationError):
            parse_bundle("100,0,0,0,0,0")
        self.assertEqual(parse_bundle("100,0,0,0,0,0", delimiter="comma")["channels"]["C"], "100")

    def test_ambiguous_or_nonliteral_inputs_rejected(self):
        for value in (100, True, "100/0;0/0/0/0", "1e2/0/0/0/0/0", "１００/0/0/0/0/0",
                      "=CONCAT(A1:A6)", "NaN/0/0/0/0/0", "101/0/0/0/0/-1"):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                parse_bundle(value)

    def test_recipe_validation_rejects_boolean_nonfinite_and_extreme_precision(self):
        for value in (True, "NaN", "Infinity", "1e-65"):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                validate_recipe([value, 0, 0, 0, 0, 100])

    def test_strict_json_refuses_duplicate_and_nonfinite_values(self):
        for value in ('{"x":1,"x":2}', '{"x":NaN}', '{"x":Infinity}', '[' * 66 + '0' + ']' * 66):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                strict_json(value)


class WorkbookSafeguards(TempTest):
    def workbook(self, cells, merge=None):
        wb = Workbook()
        ws = wb.active
        ws.title = "Samples"
        for coordinate, value in cells.items():
            ws[coordinate] = value
        if merge:
            ws.merge_cells(merge)
        path = self.root / "synthetic.xlsx"
        wb.save(path)
        wb.close()
        return path

    def test_whole_blank_rows_columns_and_shifted_range_are_preserved(self):
        path = self.workbook({"B2": "100/0/0/0/0/0", "D4": "0/100/0/0/0/0"})
        grid = read_workbook(path, sheet="Samples", cell_range="B2:D4")
        self.assertEqual((grid["rows"], grid["columns"]), (3, 3))
        self.assertEqual(len(grid["cells"]), 9)
        self.assertEqual([(c["grid_row"], c["grid_column"]) for c in grid["cells"] if c["role"] == "sample"], [(0, 0), (2, 2)])
        self.assertEqual(sum(c["role"] == "hole" for c in grid["cells"]), 7)

    def test_invalid_scalar_is_not_a_hole(self):
        path = self.workbook({"B2": 100, "C2": "100/0/0/0/0/0"})
        grid = read_workbook(path, cell_range="B2:C2")
        self.assertEqual(grid["cells"][0]["role"], "invalid")
        self.assertEqual(grid["errors"][0]["cell"], "B2")

    def test_merged_cells_intersecting_grid_rejected(self):
        path = self.workbook({"B2": "100/0/0/0/0/0"}, merge="B2:C2")
        with self.assertRaises(ValidationError):
            read_workbook(path, cell_range="B2:D2")

    def test_entity_xml_rejected_before_workbook_loading(self):
        path = self.workbook({"A1": "100/0/0/0/0/0"})
        with zipfile.ZipFile(path) as z:
            entries = [(i.filename, z.read(i)) for i in z.infolist()]
        hostile = b'<!DOCTYPE worksheet [<!ENTITY xxe SYSTEM "file:///DO_NOT_READ">]><worksheet>&xxe;</worksheet>'
        bad = self.archive([(n, hostile if n == "xl/worksheets/sheet1.xml" else b) for n, b in entries], "hostile.xlsx")
        with self.assertRaisesRegex(ValidationError, "DTD/entities|unsafe"):
            read_workbook(bad)


class ArchiveSafeguards(TempTest):
    def test_archive_fixture_preserves_requested_raw_names_on_windows(self):
        with patch.object(zipfile.os,'sep','\\'):
            path=self.archive([('project.json','{}'),('x\\y','x')])
            with zipfile.ZipFile(path) as archive:
                self.assertEqual(archive.infolist()[1].orig_filename,'x\\y')
    def test_unsafe_windows_and_traversal_names_rejected(self):
        for name in ("../outside", "/absolute", "x\\y", "C:drive", "a//b", "a/./b", "NUL.txt", "a/COM1", "trailing.", "a?b"):
            with self.subTest(name=name):
                path = self.archive([("project.json", "{}"), (name, "x")])
                with self.assertRaises(ValidationError):
                    archive_inventory(path)

    def test_duplicate_and_case_colliding_entries_rejected(self):
        for names in (("a", "a"), ("Mesh.stl", "mesh.STL"), ("caf\u00e9", "cafe\u0301")):
            with self.subTest(names=names):
                path = self.archive([("project.json", "{}"), (names[0], "a"), (names[1], "b")])
                with self.assertRaises(ValidationError):
                    archive_inventory(path)

    def test_symlink_entry_rejected(self):
        entry = zipfile.ZipInfo("link")
        entry.create_system = 3
        entry.external_attr = (stat.S_IFLNK | 0o777) << 16
        with self.assertRaises(ValidationError):
            archive_inventory(self.archive([("project.json", "{}"), (entry, "outside")]))

    def test_entry_count_and_expansion_limits_enforced(self):
        path = self.archive([("project.json", "{}"), ("blob", "payload")])
        for key, limit in (("MAX_ENTRIES", 1), ("MAX_ENTRY", 3), ("MAX_TOTAL", 4), ("MAX_JSON", 1)):
            with self.subTest(limit=key), patch.object(core, key, limit), self.assertRaises(ValidationError):
                archive_inventory(path)

    def test_required_root_project_and_strict_json(self):
        for entries in ([('folder/project.json', '{}')], [('project.json', '{"x":1,"x":2}')]):
            with self.subTest(entries=entries), self.assertRaises(ValidationError):
                archive_inventory(self.archive(entries))

    def test_duplicate_occurrences_and_dangling_bodies_rejected(self):
        doc = archive_inventory(SOURCE)["project_document"]
        duplicate = copy.deepcopy(doc)
        duplicate["project"]["rootOccurrences"].append(copy.deepcopy(duplicate["project"]["rootOccurrences"][0]))
        dangling = copy.deepcopy(doc)
        dangling["project"]["trays"][0]["trayBodies"][0]["occurrence"] = "not-present"
        for altered in (duplicate, dangling):
            with self.assertRaises(ValidationError):
                project_index(altered)


class MaterialSafeguards(unittest.TestCase):
    def setUp(self):
        self.provider = EcoFlexMaterialProvider()

    def test_six_channels_and_exact_fractional_alpha(self):
        recipe = parse_bundle("10/20/30/0/10.125/29.875")
        material = self.provider.sample(recipe)
        metadata = material["metadata"]
        self.assertEqual(metadata["cmykwt_dm_clear"], Decimal("29.875"))
        self.assertEqual(material["threeDAppearance"]["a"], Decimal("225.125"))
        self.assertEqual(metadata["cmykwt_dm_used_materials_id"], [11934, 11932, 11933, 11935, 11931])
        self.assertEqual(strict_json(exact_json(material)), material)

    def test_each_pure_channel_uses_observed_resin_id(self):
        for channel, resin_id in CHANNEL_RESIN_IDS.items():
            values = ["100" if c == channel else "0" for c in "CMYKWT"]
            with self.subTest(channel=channel):
                material = self.provider.sample(validate_recipe(values))
                self.assertEqual(material["metadata"]["cmykwt_dm_used_materials_id"], [int(resin_id)])

    def test_direct_label_black_is_not_k100_mixer(self):
        direct = self.provider.direct_black()
        mixer = self.provider.sample(parse_bundle("0/0/0/100/0/0"))
        self.assertEqual(direct["vendorMaterialId"], "VEcoFlxBlk")
        self.assertEqual(direct["metadata"]["resinId"], 11936)
        self.assertEqual(direct["metadata"]["materialType"], "base_material")
        self.assertEqual(mixer["metadata"]["materialType"], "cmykwt_dm_mixer")
        self.assertNotIn("cmykwt_dm_cyan", direct["metadata"])

    def test_precision_loss_and_tolerance_refused_not_rounded(self):
        for recipe in (parse_bundle("0.1234567890123456789/0/0/0/0/99.8765432109876543211"),
                       parse_bundle("99.99/0/0/0/0/0", tolerance=Decimal(".01"))):
            with self.assertRaises(ValidationError):
                self.provider.sample(recipe)

    def test_floats_booleans_and_forged_totals_refused(self):
        for key, value in (("C", 100.0), ("C", True), ("total", "99"), ("sum_tolerance", ".1")):
            recipe = parse_bundle("100/0/0/0/0/0")
            if key in "CMYKWT":
                recipe["channels"][key] = value
            else:
                recipe[key] = value
            with self.subTest(key=key), self.assertRaises(ValidationError):
                self.provider.sample(recipe)

    def test_fresh_material_has_no_aliasing_or_picker_state(self):
        recipe = parse_bundle("10/20/30/0/10/30")
        a, b = self.provider.sample(recipe), self.provider.sample(recipe)
        self.assertNotEqual(a["id"], b["id"])
        self.assertNotIn("pickerSelectedColor", a["metadata"])
        a["metadata"]["cmykwt_dm_cyan"] = 77
        self.assertEqual(b["metadata"]["cmykwt_dm_cyan"], 10)
        self.assertEqual(recipe["channels"]["C"], "10")

    def test_provider_profile_mutation_and_print_claims(self):
        self.assertFalse(self.provider.evidence()["print_ready"])
        self.provider.channel_resin_ids["C"] = "11936"
        with self.assertRaises(ValidationError):
            self.provider.sample(parse_bundle("100/0/0/0/0/0"))
        with self.assertRaises(ValidationError):
            EcoFlexMaterialProvider(profile="unknown")


class GeometryAndCandidateSafeguards(TempTest):
    def setUp(self):
        super().setUp()
        self.grid = read_workbook(WORKBOOK, sheet="Samples", cell_range="B2:E3")
        self.proposal = propose_grid_mapping(self.grid, SOURCE, selection=copy.deepcopy(SELECTION))
        self.candidate_id = self.proposal["orientation_candidates"][0]["candidate_id"]

    def prepared(self, source=SOURCE):
        return mapping_from_geometry_review(self.grid, source, selection=copy.deepcopy(SELECTION),
                                            candidate_id=self.candidate_id, reviewed=True)

    def write(self, source=SOURCE, name="candidate.PRINT", **overrides):
        options = dict(selection=copy.deepcopy(SELECTION), candidate_id=self.candidate_id,
                       assignment_policy={"mode": "overwrite_targets"}, reviewed=True)
        options.update(overrides)
        return write_reviewed_candidate_copy(self.grid, source, self.root / name, **options)

    def test_synthetic_fixture_hashes(self):
        self.assertEqual(file_hash(SOURCE), "76973f60dc4a77e698e074c39ecf78bfe16253302235e122ad076a6c4e04dff0")
        self.assertEqual(file_hash(WORKBOOK), "3efba8b6aaec3098d30351bf7708631e1a0671c49c9c464b4cf3ede239e50c21")

    def test_candidate_flush_uses_writable_descriptor_for_windows(self):
        real_open=builtins.open;real_fsync=os.fsync;opened=[]
        def track(*args,**kwargs):
            handle=real_open(*args,**kwargs);opened.append(handle);return handle
        def flush(descriptor):
            handle=next((h for h in reversed(opened) if not h.closed and h.fileno()==descriptor),None)
            self.assertIsNotNone(handle)
            self.assertTrue(handle.writable(),'Windows fsync requires a write-capable descriptor.')
            return real_fsync(descriptor)
        with patch('builtins.open',side_effect=track),patch('gridcad.candidate_writer.os.fsync',side_effect=flush):
            self.write()

    def test_unique_orientation_still_needs_explicit_review(self):
        self.assertEqual(len(self.proposal["orientation_candidates"]), 1)
        self.assertFalse(self.proposal["mapping_confirmed"])
        self.assertFalse(self.proposal["writer_enabled"])
        with self.assertRaises(ValidationError):
            self.write(reviewed=False)
        self.assertFalse((self.root / "candidate.PRINT").exists())

    def test_holes_and_explicit_label_roles_are_not_compressed(self):
        candidate = self.proposal["orientation_candidates"][0]
        self.assertEqual(candidate["transform"], IDENTITY)
        self.assertEqual(candidate["holes"], [{"source_grid": [0, 2], "target_grid": [0, 2], "occurrence_ids": []}])
        self.assertEqual(len(candidate["assignments"]), 7)
        self.assertEqual(self.proposal["selected_label_counts"], {"column_label": 8, "row_label": 4})
        self.assertEqual(candidate["label_text_association"], "not_inferred")

    def test_all_eight_transforms_are_bijective(self):
        for transpose, rows, columns in itertools.product((False, True), repeat=3):
            transform = dict(IDENTITY, transpose=transpose, mirror_rows=rows, mirror_columns=columns,
                             row_offset=4, column_offset=6)
            actual = {transform_position(r, c, 2, 3, transform) for r in range(2) for c in range(3)}
            expected = {(r + 4, c + 6) for r in range(3 if transpose else 2) for c in range(2 if transpose else 3)}
            self.assertEqual(actual, expected)

    def test_transform_requires_explicit_typed_fields(self):
        for transform in ({}, dict(IDENTITY, transpose=1), dict(IDENTITY, row_offset=True)):
            with self.assertRaises(ValidationError):
                transform_position(0, 0, 2, 3, transform)

    def test_label_roles_cannot_overlap_samples_or_be_omitted(self):
        overlapping = copy.deepcopy(SELECTION)
        overlapping["role_regions_mm"]["row_label"] = SELECTION["region_mm"]
        with self.assertRaises(ValidationError):
            propose_grid_mapping(self.grid, SOURCE, selection=overlapping)
        incomplete = {"region_mm": SELECTION["region_mm"]}
        proposal = propose_grid_mapping(self.grid, SOURCE, selection=incomplete)
        with self.assertRaises(ValidationError):
            self.write(selection=incomplete, candidate_id=proposal["orientation_candidates"][0]["candidate_id"])

    def test_region_cutting_body_is_rejected(self):
        selection = copy.deepcopy(SELECTION)
        sample = next(body for body in self.proposal["geometry_bodies"]
                      if body["occurrence_id"] == "case_profile-sample-r0-c0")
        lo, hi = sample["bounds_mm"]
        selection["region_mm"][0] = (lo[0] + hi[0]) / 2
        with self.assertRaises(ValidationError):
            propose_grid_mapping(self.grid, SOURCE, selection=selection)

    def test_stale_candidate_and_swapped_mapping_rejected(self):
        with self.assertRaises(ValidationError):
            self.write(candidate_id="stale")
        prepared = self.prepared()
        samples = [t for t in prepared["mapping"]["targets"] if t["role"] == "sample"]
        samples[0]["occurrence_ids"], samples[1]["occurrence_ids"] = samples[1]["occurrence_ids"], samples[0]["occurrence_ids"]
        with self.assertRaises(ValidationError):
            write_candidate_copy(self.grid, SOURCE, prepared["mapping"], prepared["witness"],
                                 EcoFlexMaterialProvider(), self.root / "swapped.PRINT", assignment_policy={"mode": "overwrite_targets"})
        self.assertFalse((self.root / "swapped.PRINT").exists())

    def test_blank_cannot_be_assigned_as_a_sample(self):
        prepared = self.prepared()
        altered = copy.deepcopy(self.grid)
        altered["cells"][0].update(role="hole", recipe=None)
        preview = mapping_preview(altered, SOURCE, prepared["mapping"])
        self.assertFalse(preview["input_validation_passed"])
        self.assertTrue(any("Blank grid hole" in issue["error"] for issue in preview["issues"]))

    def test_unknown_units_and_nonidentity_occurrence_transform_rejected(self):
        for key, value in (("units", "UNKNOWN"), ("transform", [1, 0, 0, Decimal(".1"), 0, 1, 0, 0, 0, 0, 1, 0, 0, 0, 0, 1])):
            source = self.mutate_project(lambda doc: doc["project"]["rootOccurrences"][0].__setitem__(key, value))
            with self.subTest(key=key), self.assertRaises(ValidationError):
                inspect_geometry(source)

    def test_stl_degenerate_nonfinite_and_malformed_refused(self):
        for data in (b"not stl", b"solid x\nendsolid x", b"solid x\nfacet normal 0 0 1\nouter loop\nvertex 0 0 0\nvertex 0 0 0\nvertex 0 0 0\nendloop\nendfacet\nendsolid x",
                     b"solid x\nfacet normal nan 0 1\nouter loop\nvertex 0 0 0\nvertex 1 0 0\nvertex 0 1 0\nendloop\nendfacet\nendsolid x"):
            with self.assertRaises(ValidationError):
                read_stl(data)

    def test_unknown_profile_and_schema_stay_blocked(self):
        for mutate in (lambda doc: doc["version"]["versionNumbers"]["print"].__setitem__("internal", 9999),
                       lambda doc: doc["project"]["gender"].__setitem__("model", "unknown")):
            source = self.mutate_project(mutate)
            self.assertFalse(candidate_readiness(source)["candidate_allowed"])
            with self.assertRaises(ValidationError):
                self.write(source=source)

    def test_candidate_preserves_input_payloads_and_nonmaterial_state(self):
        original_hash = file_hash(SOURCE)
        result = self.write()
        output = self.root / "candidate.PRINT"
        self.assertEqual(file_hash(SOURCE), original_hash)
        self.assertEqual(result["status"], "UNVERIFIED_CANDIDATE_COPY")
        for key in ("print_ready", "gui_validated", "production_write", "vendor_loader_reopen_validated"):
            self.assertFalse(result[key])
        audit = audit_projects(SOURCE, output)
        for key in ("added_archive_entries", "removed_archive_entries", "changed_nonproject_payloads"):
            self.assertEqual(audit[key], [])
        before = archive_inventory(SOURCE)["project_document"]
        after = archive_inventory(output)["project_document"]
        for doc in (before, after):
            _, nodes, bodies = project_index(doc)
            for oid in result["changed_occurrence_ids"]:
                nodes[oid].pop("material")
                bodies[oid][0][1].pop("material")
        self.assertEqual(semantic_differences(before, after), [])
        self.assertEqual(result["changed_objects"] + result["skipped_exact_matches"], 19)
        self.assertGreater(result["changed_objects"], 0)

    def test_no_overwrite_and_no_input_as_output(self):
        existing = self.root / "candidate.PRINT"
        existing.write_bytes(b"do not replace")
        with self.assertRaises(ValidationError):
            self.write()
        self.assertEqual(existing.read_bytes(), b"do not replace")
        prepared = self.prepared()
        original_hash = file_hash(SOURCE)
        with self.assertRaises(ValidationError):
            write_candidate_copy(self.grid, SOURCE, prepared["mapping"], prepared["witness"], EcoFlexMaterialProvider(),
                                 SOURCE, assignment_policy={"mode": "overwrite_targets"})
        self.assertEqual(file_hash(SOURCE), original_hash)

    def test_assignment_policy_required_and_fake_completed_rejected(self):
        samples = [a["occurrence_id"] for a in self.proposal["orientation_candidates"][0]["assignments"]]
        for policy in (None, {"mode": "fill_uncompleted", "completed_occurrence_ids": samples, "pending_occurrence_ids": []},
                       {"mode": "fill_uncompleted", "completed_occurrence_ids": [], "pending_occurrence_ids": samples[:-1]}):
            with self.subTest(policy=policy), self.assertRaises(ValidationError):
                self.write(assignment_policy=policy)
            self.assertFalse((self.root / "candidate.PRINT").exists())

    def test_exact_repeat_skips_all_materials_but_requires_fresh_review(self):
        self.write()
        first = self.root / "candidate.PRINT"
        with self.assertRaises(ValidationError):
            self.write(source=first, name="stale-review.PRINT")
        proposal = propose_grid_mapping(self.grid, first, selection=copy.deepcopy(SELECTION))
        result = self.write(source=first, name="second.PRINT", candidate_id=proposal["orientation_candidates"][0]["candidate_id"])
        self.assertEqual(result["changed_objects"], 0)
        self.assertEqual(result["skipped_exact_matches"], 19)


if __name__ == "__main__":
    unittest.main()
